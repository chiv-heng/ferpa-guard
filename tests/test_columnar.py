"""Bounded column evidence contracts, using generated synthetic inputs only."""
import csv
import io
import json
import os
from pathlib import Path
import random
import re
import subprocess
import sys
import tempfile
import unittest
from unittest import mock
import zipfile

ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT))
from shared import pii_engine as e

def bound_counts(evidence):
    return {name: item.bound_count for name, item in evidence.items()}


def findings(si, **kwargs):
    return e.scan_content(si.content,si.header_line_indices,column_evidence=si.column_evidence,**kwargs)


class DelimitedTests(unittest.TestCase):
    def read(self,text,extension='.csv'):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/('synthetic'+extension);p.write_text(text)
            return e.read_file_content(str(p))

    def test_every_specific_alias_binds(self):
        aliases={'SASID':['sasid','state_studentnumber'], 'STUDENT_ID_LABELED':['student_id','pupil_id','sis_id','ps_id','student_number','studentNumber'], 'LUNCH_PIN':['lunch_pin','meal_pin','cafeteria_pin']}
        for pattern,names in aliases.items():
            width=10 if pattern=='SASID' else 4
            for name in names:
                si=self.read(name+'\n'+'0'*width+'\n')
                with self.subTest(pattern=pattern,name=name):
                    self.assertEqual(bound_counts(si.column_evidence),{pattern:1})
                    fs=findings(si)
                    self.assertTrue(next(f['column_bound'] for f in fs if f['pattern_name']==pattern))

    def test_generic_alias_requires_independent_student_header(self):
        for name in ['dcid','personID','stateID']:
            width=10 if name=='stateID' else 4
            alone=self.read(name+',staff_name\n'+'0'*width+',SYNTH\n')
            student=self.read(name+',student_number\n'+'0'*width+',\n')
            with self.subTest(name=name):
                self.assertEqual(alone.column_evidence,{})
                self.assertTrue(student.column_evidence)

    def test_loose_context_does_not_bind_staff(self):
        si=self.read('personID,school,teacher,record\n'+'0'*4+',SYNTH,SYNTH,SYNTH\n')
        self.assertEqual(si.column_evidence,{})

    def test_normalization_is_exact_and_ascii(self):
        self.assertTrue(self.read('  StUdEnT-ID  \n'+'0'*7).column_evidence)
        for alias in ['prefix_student_id','student.studentNumber','[Students]DCID','lunch_id','ſtudent_id']:
            with self.subTest(alias=alias):
                self.assertEqual(self.read(alias+'\n'+'0'*7).column_evidence,{})

    def test_tsv_bom_quotes_and_embedded_newlines(self):
        si=self.read('\ufeff"student id"\t"notes"\r\n"'+'0'*7+'"\t"SYNTH\nSYNTH \"\"quoted\"\""\r\n','.tsv')
        self.assertEqual(bound_counts(si.column_evidence),{'STUDENT_ID_LABELED':1})
        self.assertFalse(si.reader_error)
        self.assertEqual(si.header_line_indices,set())

    def test_headers_placeholders_and_invalid_widths_are_not_values(self):
        for cell in ['', 'NULL','N/A','REDACTED','***','0'*3,'0'*11,'x'+'0'*7,'0'*7+'.0']:
            with self.subTest(kind=len(cell)):
                self.assertEqual(self.read('student_id\n"'+cell+'"\n').column_evidence,{})
        self.assertEqual(self.read('student_id\n').column_evidence,{})

    def test_headerless_and_prose_extensions_do_not_bind(self):
        self.assertEqual(self.read('0000,0000\n0000,0000\n').column_evidence,{})
        for ext in ['.txt','.json','.md','.xls']:
            self.assertEqual(self.read('student_id\n0000\n',ext).column_evidence,{})

    def test_duplicate_columns_stay_indexed(self):
        self.assertEqual(bound_counts(self.read('student_id,student_id\n0000,0000\n').column_evidence),{'STUDENT_ID_LABELED':2})

    def test_ragged_and_malformed_rows_hold_and_discard_evidence(self):
        for text in ['student_id,note\n0000,ok\n0000\n','student_id\n0000\n"unterminated','student_id\n0000\n"0000"bad\n','student_id\n0000\n\ufffd\n']:
            si=self.read(text)
            with self.subTest(length=len(text)):
                self.assertEqual(si.column_evidence,{})
                self.assertEqual(si.reader_error,'EXTRACTION_FAILED')
                self.assertTrue(si.content)

    def test_skip_patterns_applies_to_both_routes(self):
        fs=findings(self.read('student_id\n0000\n'),skip_patterns={'STUDENT_ID_LABELED'})
        self.assertNotIn('STUDENT_ID_LABELED',{f['pattern_name'] for f in fs})

    def test_bound_and_inline_merge_without_cooccurrence_inflation(self):
        si=self.read('student_id,note\n0000,student_id=0000\n')
        fs=findings(si)
        found=[f for f in fs if f['pattern_name']=='STUDENT_ID_LABELED']
        self.assertEqual(len(found),1);self.assertEqual(found[0]['count'],2)

    def test_lunch_scoring_matches_inline_not_hard_floor(self):
        si=self.read('lunch_pin\n0000\n')
        bound=findings(si);inline=e.scan_content('lunch_pin=0000\n')
        self.assertEqual(e.worst_action(bound),e.worst_action(inline))
        self.assertEqual(bound[0]['confidence'],'medium')

    def test_early_exit_keeps_action_parity(self):
        si=self.read('student_id,email\n0000,synthetic@example.invalid\n')
        self.assertEqual(e.worst_action(findings(si,early_exit=True)),e.worst_action(findings(si)))

    def test_extensionless_symlink_uses_target_reader(self):
        with tempfile.TemporaryDirectory() as tmp:
            target=Path(tmp)/'synthetic.csv';target.write_text('student_id\n0000\n')
            alias=Path(tmp)/'alias';alias.symlink_to(target)
            self.assertEqual(bound_counts(e.read_file_content(str(alias)).column_evidence),{'STUDENT_ID_LABELED':1})


class CSVAllocationTests(unittest.TestCase):
    def module(self):
        from shared import columnar
        return columnar

    def test_excess_columns_never_reach_parser_or_copy(self):
        c=self.module()
        with mock.patch.object(c,'_parse_record') as parser,mock.patch.object(c,'_copy_record') as copy:
            with self.assertRaises(c.ColumnarError): c.csv_evidence(','*65537+'\n',',')
            parser.assert_not_called();copy.assert_not_called()

    def test_excess_data_record_stops_before_allocation(self):
        c=self.module();real=c._parse_record
        with mock.patch.object(c,'_parse_record',wraps=real) as parser:
            with self.assertRaises(c.ColumnarError): c.csv_evidence('student_id\n'+','*1_000_000,',')
            self.assertEqual(parser.call_count,1)

    def test_exact_and_one_over_header_syntax_limit(self):
        c=self.module()
        with mock.patch.object(c,'MAX_HEADER_CHARS',8):
            self.assertEqual(c.csv_evidence('"abcdef"\n',','),{})
            with mock.patch.object(c,'_parse_record') as parser:
                with self.assertRaises(c.ColumnarError): c.csv_evidence('"abcdefg"\n',',')
                parser.assert_not_called()

    def test_decoded_quotes_and_field_limit_without_global_mutation(self):
        c=self.module();before=csv.field_size_limit()
        with mock.patch.object(c,'MAX_FIELD_CHARS',4):
            self.assertEqual(c.csv_evidence('""""""""""\n',','),{})
            with self.assertRaises(c.ColumnarError): c.csv_evidence('""""""""""""\n',',')
        self.assertEqual(csv.field_size_limit(),before)

    def test_blank_records_and_record_budget(self):
        c=self.module()
        with mock.patch.object(c,'MAX_RECORDS',2):
            self.assertEqual(c.csv_evidence('\n\n',','),{})
            with self.assertRaises(c.ColumnarError): c.csv_evidence('\n\n\n',',')
        self.assertEqual(c.csv_evidence('\ufeff',','),{})

    def test_quoted_delimiters_do_not_count_as_columns(self):
        c=self.module()
        self.assertEqual(c.csv_evidence('"'+','*6000+'"\n',','),{})

    def test_authoritative_parser_disagreement_holds(self):
        c=self.module()
        with mock.patch.object(c,'_parse_record',return_value=['WRONG','WIDTH']):
            with self.assertRaises(c.ColumnarError): c.csv_evidence('one\n',',')

    def test_differential_synthetic_csv_dialects(self):
        c=self.module();rng=random.Random(121)
        for delimiter in [',','\t']:
            for _ in range(80):
                rows=[[''.join(rng.choice('ab,\t\r\n" ') for _ in range(rng.randrange(20))) for _ in range(4)] for _ in range(3)]
                buf=io.StringIO(newline='');writer=csv.writer(buf,delimiter=delimiter);writer.writerows(rows)
                self.assertEqual(list(c.validated_rows(buf.getvalue(),delimiter)),rows)
        for text in ['a"b,c\n','"a\r\nb",c\r\n','a,b\r','"a"bad\n','"unterminated']:
            try: expected=list(csv.reader(io.StringIO(text,newline=''),strict=True))
            except csv.Error:
                with self.assertRaises(c.ColumnarError): list(c.validated_rows(text,','))
            else: self.assertEqual(list(c.validated_rows(text,',')),expected)


class WorkbookBindingTests(unittest.TestCase):
    def make(self,root,rows,other=None):
        import openpyxl
        p=root/'synthetic.xlsx';w=openpyxl.Workbook();ws=w.active
        for row in rows:ws.append(row)
        if other:
            ws=w.create_sheet('other')
            for row in other:ws.append(row)
        w.save(p);w.close();return p

    def rewrite(self,p,change):
        with zipfile.ZipFile(p) as z: parts={n:z.read(n) for n in z.namelist()}
        parts['xl/worksheets/sheet1.xml']=change(parts['xl/worksheets/sheet1.xml'].decode()).encode()
        with zipfile.ZipFile(p,'w') as z:
            for n,b in parts.items():z.writestr(n,b)

    def test_blank_positions_and_numeric_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=self.make(Path(tmp),[['note',None,'student_id'],[None,'SYNTH',4000],[None,None,'0000']])
            self.assertEqual(bound_counts(e.read_file_content(str(p)).column_evidence),{'STUDENT_ID_LABELED':2})

    def test_sheet_context_does_not_bleed(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=self.make(Path(tmp),[['student_number'],['NULL']],other=[['personID'],[4000]])
            self.assertEqual(bound_counts(e.read_file_content(str(p)).column_evidence),{})

    def test_wrong_or_absent_dimensions_do_not_hide_values(self):
        for dimension in ['<dimension ref="A1:A1"/>','']:
            with tempfile.TemporaryDirectory() as tmp:
                p=self.make(Path(tmp),[['note','student_id'],['SYNTH','0000']])
                self.rewrite(p,lambda text:re.sub(r'<dimension[^>]*/>',dimension,text))
                self.assertEqual(bound_counts(e.read_file_content(str(p)).column_evidence),{'STUDENT_ID_LABELED':1})

    def test_excess_dimension_rejected_before_row_production(self):
        from shared import columnar
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        with tempfile.TemporaryDirectory() as tmp:
            p=self.make(Path(tmp),[['student_id'],['0000']])
            self.rewrite(p,lambda text:re.sub(r'<dimension[^>]*/>','<dimension ref="A1:ZZZ2"/>',text))
            with mock.patch.object(ReadOnlyWorksheet,'iter_rows',side_effect=AssertionError('producer called')) as producer:
                si=e.read_file_content(str(p))
            producer.assert_not_called();self.assertEqual(si.reader_error,'EXTRACTION_FAILED')

    def test_observed_outside_dimension_rejected_before_production(self):
        from shared import columnar
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        with tempfile.TemporaryDirectory() as tmp:
            p=self.make(Path(tmp),[['student_id'],['0000']])
            self.rewrite(p,lambda text:text.replace('r="A2"','r="ZZZ2"'))
            with mock.patch.object(ReadOnlyWorksheet,'iter_rows') as producer:
                si=e.read_file_content(str(p))
            producer.assert_not_called();self.assertEqual(si.reader_error,'EXTRACTION_FAILED')

    def test_explicit_bounded_producer_arguments(self):
        from shared import columnar
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        original=ReadOnlyWorksheet.iter_rows;calls=[]
        def checked(ws,*args,**kwargs):
            calls.append(kwargs);return original(ws,*args,**kwargs)
        with tempfile.TemporaryDirectory() as tmp:
            p=self.make(Path(tmp),[['student_id'],['0000']])
            with mock.patch.object(ReadOnlyWorksheet,'iter_rows',checked):si=e.read_file_content(str(p))
            self.assertTrue(si.column_evidence)
        self.assertEqual(calls,[dict(min_row=1,min_col=1,max_row=2,max_col=1,values_only=True)])

    def test_boolean_fraction_formula_are_not_ids(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=self.make(Path(tmp),[['student_id'],[True],[4000.5],['=4000'],[4000.0]])
            self.assertEqual(bound_counts(e.read_file_content(str(p)).column_evidence),{'STUDENT_ID_LABELED':1})

    def test_later_bad_sheet_retains_hold_after_critical(self):
        from shared import columnar as c
        with tempfile.TemporaryDirectory() as tmp:
            p=self.make(Path(tmp),[['student_id'],['0000']],other=[['student_id'],['0000']])
            with mock.patch.object(c,'MAX_VISITED_CELLS',3):si=e.read_file_content(str(p))
            self.assertEqual(si.reader_error,'EXTRACTION_FAILED');self.assertEqual(si.column_evidence,{})


class XMLPreflightControls(unittest.TestCase):
    def run_xml(self, body, dimension='', budget=None):
        from types import SimpleNamespace
        from shared import columnar as c
        source=io.BytesIO(('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'+dimension+'<sheetData>'+body+'</sheetData></worksheet>').encode())
        workbook=SimpleNamespace(_archive=SimpleNamespace(open=lambda path:source))
        sheet=SimpleNamespace(_worksheet_path='synthetic.xml')
        try:return c.xlsx_bounds(workbook,sheet,{} if budget is None else budget)
        finally:self.assertTrue(source.closed)

    def test_malformed_noncanonical_and_duplicate_coordinates_hold(self):
        from shared import columnar as c
        bad=['<row r="01"/>','<row r="1.0"/>','<row r="1000001"/>',
             '<row r="1"/><row r="1"/>','<row r="2"/><row r="1"/>',
             '<row r="1"><c r="a1"/></row>','<row r="1"><c r="A2"/></row>',
             '<row r="1"><c r="A1"/><c r="A1"/></row>',
             '<row r="1"><foreign/></row>','<row r="1"><c r="AAAA1"/></row>']
        for body in bad:
            with self.assertRaises(c.ColumnarError):self.run_xml(body)

    def test_exact_rectangular_cell_budget_and_one_over(self):
        from shared import columnar as c
        with mock.patch.object(c,'MAX_VISITED_CELLS',6):
            self.assertEqual(self.run_xml('<row r="2"><c r="C2"/></row>'),(2,3))
            with self.assertRaises(c.ColumnarError):self.run_xml('<row r="2"><c r="D2"/></row>')
            with self.assertRaises(c.ColumnarError):self.run_xml('<row r="2"><c r="C2"/></row>',budget={'cells':1})

    def test_implicit_coordinates_and_declared_rectangle_preserved(self):
        self.assertEqual(self.run_xml('<row><c/><c/></row><row><c/></row>'),(2,2))
        self.assertEqual(self.run_xml('<row><c/></row>','<dimension ref="A1:D3"/>'),(3,4))

    def test_unknown_private_seam_fails_closed(self):
        from shared import columnar as c
        with self.assertRaises(c.ColumnarError):c.xlsx_bounds(object(),object(),{})

    def test_dtd_and_entity_are_refused(self):
        from types import SimpleNamespace
        from shared import columnar as c
        source=io.BytesIO(b'<!DOCTYPE worksheet [<!ENTITY synthetic "synthetic">]><worksheet/>')
        workbook=SimpleNamespace(_archive=SimpleNamespace(open=lambda path:source))
        with self.assertRaises(c.ColumnarError):c.xlsx_bounds(workbook,SimpleNamespace(_worksheet_path='synthetic.xml'),{})
        self.assertTrue(source.closed)


if __name__=='__main__':unittest.main()
