#!/usr/bin/env python3
"""Contract tests for fail-closed file readers (blind spot 3).

All fixtures are synthetic.  New API names are deliberately resolved inside
test methods so this module still imports and reaches a unittest summary when
run against the pre-fix engine.
"""

import json
import os
import subprocess
import sys
import tempfile
import time
import types
import unittest
from pathlib import Path
from unittest import mock

# Match tests/test_pii_guardian.py: add the project root for shared imports.
_project_root = Path(__file__).parent.parent
sys.path.insert(0, str(_project_root))

import shared.pii_engine as pii_engine
from shared.pii_engine import ScanInput, read_file_content, scan_content, worst_action


HOOK_SCRIPT = _project_root / "claude-code" / "pii_guardian.py"
REPORT_SCRIPT = _project_root / "claude-code" / "pii_scan_report.py"
RAW_EXCEPTION_SENTINEL = "RAW_READER_EXCEPTION_DO_NOT_LEAK_8675309"


def _run_script(script, args, *, stdin=None, home, env_extra=None):
    env = os.environ.copy()
    env["HOME"] = str(home)
    if env_extra:
        env.update(env_extra)
    return subprocess.run(
        [sys.executable, str(script), *args],
        input=stdin,
        capture_output=True,
        text=True,
        env=env,
        timeout=30,
    )


def _run_hook(filepath, *, home, env_extra=None):
    payload = json.dumps({"tool_name": "Read", "tool_input": {"file_path": str(filepath)}})
    return _run_script(HOOK_SCRIPT, [], stdin=payload, home=home, env_extra=env_extra)


def _shim_env(shim_dir):
    old_path = os.environ.get("PYTHONPATH", "")
    pythonpath = str(shim_dir) if not old_path else os.pathsep.join((str(shim_dir), old_path))
    return {"PYTHONPATH": pythonpath}


def _write_import_failure_shim(directory, module_name):
    (Path(directory) / f"{module_name}.py").write_text("raise ImportError('synthetic missing dependency')\n")


def _write_partial_xlsx_shim(directory):
    (Path(directory) / "openpyxl.py").write_text(
        "class Sheet:\n"
        "    def iter_rows(self, values_only=True):\n"
        "        yield ('record', 'SSN: 123-45-6789')\n"
        f"        raise RuntimeError('{RAW_EXCEPTION_SENTINEL}')\n"
        "class Workbook:\n"
        "    sheetnames = ['Synthetic']\n"
        "    def __getitem__(self, name): return Sheet()\n"
        "    def close(self): pass\n"
        "def load_workbook(*args, **kwargs): return Workbook()\n"
    )


def _audit_records(home):
    audit_path = Path(home) / ".claude" / "logs" / "ferpa-guard-audit.jsonl"
    if not audit_path.exists():
        return []
    return [json.loads(line) for line in audit_path.read_text().splitlines() if line.strip()]


class TestReaderOutcomeContract(unittest.TestCase):
    def test_missing_openpyxl_sets_code(self):
        """Spec 2.1: absent openpyxl returns MISSING_DEPENDENCY:openpyxl."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "synthetic_pii.xlsx"
            path.write_bytes(b"synthetic xlsx placeholder with SSN 123-45-6789")
            with mock.patch.dict(sys.modules, {"openpyxl": None}):
                result = read_file_content(str(path))
        self.assertIsInstance(result, ScanInput)
        self.assertEqual(result.reader_error, "MISSING_DEPENDENCY:openpyxl")

    def test_missing_pymupdf_sets_code(self):
        """Spec 2.1: absent PyMuPDF returns MISSING_DEPENDENCY:pymupdf."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "synthetic_pii.pdf"
            path.write_bytes(b"synthetic pdf placeholder with SSN 123-45-6789")
            with mock.patch.dict(sys.modules, {"fitz": None}):
                result = read_file_content(str(path))
        self.assertIsInstance(result, ScanInput)
        self.assertEqual(result.reader_error, "MISSING_DEPENDENCY:pymupdf")

    def test_missing_python_docx_sets_code(self):
        """Spec 2.1: absent python-docx returns MISSING_DEPENDENCY:python-docx."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "synthetic_pii.docx"
            path.write_bytes(b"synthetic docx placeholder with SSN 123-45-6789")
            with mock.patch.dict(sys.modules, {"docx": None}):
                result = read_file_content(str(path))
        self.assertIsInstance(result, ScanInput)
        self.assertEqual(result.reader_error, "MISSING_DEPENDENCY:python-docx")

    def test_corrupt_xlsx_sets_open_failed(self):
        """Spec 2.1: an xlsx open-time failure returns OPEN_FAILED."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "corrupt_synthetic.xlsx"
            path.write_bytes(b"this is deliberately not an OOXML archive")
            result = read_file_content(str(path))
        self.assertEqual(result.reader_error, "OPEN_FAILED")
        self.assertEqual(result.content, "")

    def test_encrypted_pdf_sets_encrypted(self):
        """Spec 2.1: an encrypted PDF is labeled ENCRYPTED before extraction."""
        try:
            import fitz
        except ImportError:
            self.skipTest("PyMuPDF is required by the frozen encrypted-PDF fixture")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "encrypted_synthetic.pdf"
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), "Synthetic record SSN: 123-45-6789")
            doc.save(
                str(path),
                encryption=fitz.PDF_ENCRYPT_AES_256,
                owner_pw="synthetic-owner",
                user_pw="synthetic-user",
            )
            doc.close()
            result = read_file_content(str(path))
        self.assertEqual(result.reader_error, "ENCRYPTED")
        self.assertEqual(result.content, "")

    def test_mid_iteration_retains_prefix_and_both_findings(self):
        """Spec 2.1-2.2: EXTRACTION_FAILED retains and scans the extracted prefix."""
        class FakeSheet:
            def iter_rows(self, values_only=True):
                yield ("record", "SSN: 123-45-6789")
                raise RuntimeError(RAW_EXCEPTION_SENTINEL)

        class FakeWorkbook:
            sheetnames = ["Synthetic"]

            def __getitem__(self, name):
                return FakeSheet()

            def close(self):
                pass

        fake_openpyxl = types.ModuleType("openpyxl")
        fake_openpyxl.load_workbook = lambda *args, **kwargs: FakeWorkbook()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "partial_synthetic.xlsx"
            path.write_bytes(b"synthetic placeholder")
            with mock.patch.dict(sys.modules, {"openpyxl": fake_openpyxl}):
                result = read_file_content(str(path))

        self.assertEqual(result.reader_error, "EXTRACTION_FAILED")
        self.assertIn("123-45-6789", result.content)
        reader_error_finding = getattr(pii_engine, "reader_error_finding")
        findings = scan_content(result.content, result.header_line_indices)
        findings.append(reader_error_finding(result.reader_error, str(path)))
        names = {finding["pattern_name"] for finding in findings}
        self.assertIn("SSN", names)
        self.assertIn("SCAN_READER_UNAVAILABLE", names)

    def test_text_over_cap_is_scan_incomplete(self):
        """Spec 2.1, 2.4: verified text omission sets TEXT_LIMIT and blocks."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "oversized_synthetic.txt"
            path.write_text("A" * 25)
            with mock.patch.object(pii_engine, "MAX_SCAN_CHARS", 24, create=True):
                result = read_file_content(str(path))

        self.assertEqual(result.truncated, "TEXT_LIMIT")
        self.assertEqual(len(result.content), 24)
        scan_incomplete_finding = getattr(pii_engine, "scan_incomplete_finding")
        finding = scan_incomplete_finding(result.truncated, str(path))
        self.assertEqual(finding["pattern_name"], "SCAN_INCOMPLETE")
        self.assertEqual(worst_action([finding]), "block")

    def test_text_exact_cap_is_complete(self):
        """Spec 2.1 and acceptance 6: exact MAX_SCAN_CHARS has no truncation finding."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "exact_cap_synthetic.txt"
            path.write_text("A" * 24)
            with mock.patch.object(pii_engine, "MAX_SCAN_CHARS", 24, create=True):
                result = read_file_content(str(path))

        self.assertEqual(result.truncated, "")
        self.assertEqual(len(result.content), 24)
        findings = scan_content(result.content, result.header_line_indices)
        if result.reader_error:
            findings.append(
                getattr(pii_engine, "reader_error_finding")(result.reader_error, str(path))
            )
        if result.truncated:
            findings.append(
                getattr(pii_engine, "scan_incomplete_finding")(result.truncated, str(path))
            )
        operational = {"SCAN_READER_UNAVAILABLE", "SCAN_INCOMPLETE"}
        self.assertTrue(operational.isdisjoint(f["pattern_name"] for f in findings))

    def test_operational_finding_shape_and_action(self):
        """Spec 2.2: operational findings are high/high VALUE findings that block."""
        reader_error_finding = getattr(pii_engine, "reader_error_finding")
        finding = reader_error_finding("OPEN_FAILED", "/synthetic/records.xlsx")
        self.assertEqual(finding["pattern_name"], "SCAN_READER_UNAVAILABLE")
        self.assertEqual(finding["severity"], "high")
        self.assertEqual(finding["confidence"], "high")
        self.assertEqual(finding["count"], 1)
        self.assertFalse(finding["header_only"])
        self.assertFalse(finding["is_metadata"])
        self.assertEqual(worst_action([finding]), "block")


class TestHookFailClosedContract(unittest.TestCase):
    def test_missing_openpyxl_hook_blocks_with_operational_formatter(self):
        """Spec 2.2-2.3: format_unscannable_reason says could not check, never contains."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            shim = root / "shim"
            home.mkdir()
            shim.mkdir()
            _write_import_failure_shim(shim, "openpyxl")
            path = root / "synthetic_pii.xlsx"
            path.write_bytes(b"synthetic xlsx placeholder with SSN 123-45-6789")
            result = _run_hook(path, home=home, env_extra=_shim_env(shim))

        self.assertEqual(result.returncode, 2)
        message = result.stdout + result.stderr
        self.assertIn("could not check", message.lower())
        self.assertNotIn("contains", message.lower())

    def test_missing_pymupdf_hook_blocks(self):
        """Spec 2.2-2.3: missing PyMuPDF yields a block-tier hook outcome."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            shim = root / "shim"
            home.mkdir()
            shim.mkdir()
            _write_import_failure_shim(shim, "fitz")
            path = root / "synthetic_pii.pdf"
            path.write_bytes(b"synthetic pdf placeholder with SSN 123-45-6789")
            result = _run_hook(path, home=home, env_extra=_shim_env(shim))

        self.assertEqual(result.returncode, 2)
        self.assertIn("could not check", (result.stdout + result.stderr).lower())

    def test_missing_python_docx_hook_blocks(self):
        """Spec 2.2-2.3: missing python-docx yields a block-tier hook outcome."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            shim = root / "shim"
            home.mkdir()
            shim.mkdir()
            _write_import_failure_shim(shim, "docx")
            path = root / "synthetic_pii.docx"
            path.write_bytes(b"synthetic docx placeholder with SSN 123-45-6789")
            result = _run_hook(path, home=home, env_extra=_shim_env(shim))

        self.assertEqual(result.returncode, 2)
        self.assertIn("could not check", (result.stdout + result.stderr).lower())

    def test_detection_only_keeps_contains_formatter(self):
        """Spec 2.3 formatter branch 2: detection-only uses the existing contains message."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            home.mkdir()
            path = root / "synthetic.csv"
            path.write_text("record,ssn\nsynthetic,123-45-6789\n")
            result = _run_hook(path, home=home)

        self.assertEqual(result.returncode, 2)
        message = result.stdout + result.stderr
        self.assertIn("contains", message.lower())
        self.assertNotIn("could not check", message.lower())

    def test_mixed_formatter_and_audit_include_detection_and_operational_names(self):
        """Spec 2.3/acceptance 9: format_partial_scan_reason audits both pattern names."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            shim = root / "shim"
            home.mkdir()
            shim.mkdir()
            _write_partial_xlsx_shim(shim)
            path = root / "partial_synthetic.xlsx"
            path.write_bytes(b"synthetic placeholder")
            result = _run_hook(path, home=home, env_extra=_shim_env(shim))
            records = _audit_records(home)

        self.assertEqual(result.returncode, 2)
        message = result.stdout + result.stderr
        self.assertIn("Social Security", message)
        self.assertIn("portion it could check", message)
        self.assertIn("could not verify the remainder", message)
        patterns = {name for record in records for name in record.get("patterns", [])}
        self.assertIn("SSN", patterns)
        self.assertIn("SCAN_READER_UNAVAILABLE", patterns)

    def test_raw_extraction_exception_never_leaks(self):
        """Spec 2.1 and acceptance 4: raw reader exceptions never enter messages or audit."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            shim = root / "shim"
            home.mkdir()
            shim.mkdir()
            _write_partial_xlsx_shim(shim)
            path = root / "partial_synthetic.xlsx"
            path.write_bytes(b"synthetic placeholder")
            result = _run_hook(path, home=home, env_extra=_shim_env(shim))
            audit_text = "\n".join(json.dumps(record) for record in _audit_records(home))

        self.assertEqual(result.returncode, 2)
        self.assertNotIn(RAW_EXCEPTION_SENTINEL, result.stdout)
        self.assertNotIn(RAW_EXCEPTION_SENTINEL, result.stderr)
        self.assertNotIn(RAW_EXCEPTION_SENTINEL, audit_text)

    def test_skip_patterns_and_skip_syntax_do_not_suppress_operational_finding(self):
        """Spec 2.2: global skip_patterns and SKIP: syntax cannot suppress operational findings."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            shim = root / "shim"
            shim.mkdir()
            _write_import_failure_shim(shim, "openpyxl")
            path = root / "synthetic.xlsx"
            path.write_bytes(b"synthetic placeholder")
            shim_env = _shim_env(shim)

            global_home = root / "global-home"
            global_home.mkdir()
            global_env = dict(shim_env)
            global_env["FERPA_GUARD_SKIP_PATTERNS"] = "SCAN_READER_UNAVAILABLE,SSN"
            global_result = _run_hook(path, home=global_home, env_extra=global_env)

            file_home = root / "file-home"
            allow_dir = file_home / ".claude"
            allow_dir.mkdir(parents=True)
            (allow_dir / "ferpa-guard-allow.txt").write_text(
                f"{path} SKIP:SCAN_READER_UNAVAILABLE,SSN\n"
            )
            file_result = _run_hook(path, home=file_home, env_extra=shim_env)

        self.assertEqual(global_result.returncode, 2)
        self.assertEqual(file_result.returncode, 2)

    def test_operational_findings_are_not_cached(self):
        """Spec 2.3 cache rule 1 and acceptance 5a: unavailable results are never cached."""
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is required to prove recovery after removing the shim")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            shim = root / "shim"
            home.mkdir()
            shim.mkdir()
            path = root / "clean_synthetic.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["color", "count"])
            wb.active.append(["blue", 3])
            wb.save(path)
            wb.close()
            before = path.stat()

            _write_import_failure_shim(shim, "openpyxl")
            first_env = _shim_env(shim)
            first_env["FERPA_GUARD_CACHE"] = "1"
            first = _run_hook(path, home=home, env_extra=first_env)
            second = _run_hook(path, home=home, env_extra={"FERPA_GUARD_CACHE": "1"})
            after = path.stat()

        self.assertEqual((before.st_mtime_ns, before.st_size), (after.st_mtime_ns, after.st_size))
        self.assertEqual(first.returncode, 2)
        self.assertEqual(second.returncode, 0)

    def test_v2_empty_cache_is_rejected_and_pii_file_rescanned(self):
        """Spec 2.3 cache rule 2 and acceptance 5b: v2 cached [] cannot allow a PII file."""
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is required for the poisoned-cache fixture")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            cache_dir = home / ".claude"
            cache_dir.mkdir(parents=True)
            path = root / "synthetic_pii.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["record", "ssn"])
            wb.active.append(["synthetic", "123-45-6789"])
            wb.save(path)
            wb.close()
            stat = path.stat()
            poisoned = {
                "version": 2,
                "entries": [{
                    "key": [str(path.resolve()), stat.st_mtime, stat.st_size],
                    "findings": [],
                    "cached_at": time.time(),
                }],
            }
            (cache_dir / "ferpa-guard-cache.json").write_text(json.dumps(poisoned))
            result = _run_hook(path, home=home, env_extra={"FERPA_GUARD_CACHE": "1"})

        self.assertEqual(result.returncode, 2)
        self.assertIn("Social Security", result.stdout + result.stderr)

    def test_allowlist_bypass_precedes_unscannable_reader(self):
        """Spec 2.3: a full allowlist bypass still works for an unscannable file."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            shim = root / "shim"
            home.mkdir()
            shim.mkdir()
            _write_import_failure_shim(shim, "openpyxl")
            path = root / "synthetic.xlsx"
            path.write_bytes(b"synthetic placeholder")
            env = _shim_env(shim)
            env["FERPA_GUARD_ALLOW"] = str(path)
            result = _run_hook(path, home=home, env_extra=env)

        self.assertEqual(result.returncode, 0)
        self.assertNotIn("could not check", (result.stdout + result.stderr).lower())


class TestScanReportFailClosedContract(unittest.TestCase):
    def test_text_report_unscannable_only_is_nonzero_and_not_clean(self):
        """Spec 2.5 and acceptance 8: text report puts unreadable-only input in unscannable and exits nonzero."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            data = root / "data"
            home.mkdir()
            data.mkdir()
            path = data / "corrupt_synthetic.xlsx"
            path.write_bytes(b"deliberately corrupt synthetic OOXML")
            result = _run_script(REPORT_SCRIPT, [str(data)], home=home)

        self.assertNotEqual(result.returncode, 0)
        output = result.stdout.lower()
        # Spec 2.5 pins the TEXT rendering as plain-language "could not check";
        # the internal bucket name "unscannable" belongs to the JSON contract
        # (asserted in the json-mode test below). Integration decision 2026-08-07:
        # the original assertion on the literal bucket name over-asserted.
        self.assertIn("could not check", output)
        self.assertRegex(output, r"could not check:\s*1")
        self.assertIn(path.name.lower(), output)
        clean_section = output.split("clean (no pii detected)", 1)
        if len(clean_section) == 2:
            self.assertNotIn(path.name.lower(), clean_section[1])

    def test_json_report_unscannable_only_is_nonzero_and_not_clean(self):
        """Spec 2.5 and acceptance 8: JSON report has unscannable, not clean, and exits nonzero."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            data = root / "data"
            home.mkdir()
            data.mkdir()
            path = data / "corrupt_synthetic.xlsx"
            path.write_bytes(b"deliberately corrupt synthetic OOXML")
            result = _run_script(REPORT_SCRIPT, [str(data), "--json"], home=home)

        self.assertNotEqual(result.returncode, 0)
        report = json.loads(result.stdout)
        self.assertIn("unscannable", report)
        self.assertEqual(report["clean"], [])
        unscannable_paths = [item.get("path", item) for item in report["unscannable"]]
        self.assertIn(str(path), unscannable_paths)
        self.assertEqual(report["summary"]["unscannable"], 1)


if __name__ == "__main__":
    unittest.main()
