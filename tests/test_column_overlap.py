"""Physical occurrence overlap contracts. All inputs are generated synthetic data."""
from array import array
import csv
import io
import importlib.util
import os
import subprocess
import time
import weakref
import json
from pathlib import Path
import random
import re
import sys
import tempfile
import unittest
from unittest import mock

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from shared import columnar as c
from shared import pii_engine as e


def patterns():
    return {name: e.PII_PATTERNS[name]['pattern'] for name in ('SASID', 'STUDENT_ID_LABELED', 'LUNCH_PIN')}


def scan(text, delimiter=','):
    evidence = c.csv_evidence(text, delimiter, patterns())
    return evidence, e.scan_content(text, column_evidence=evidence)


def counts(evidence, name):
    item = evidence[name]
    return item.bound_count, item.regex_overlap_count


class ContainmentTests(unittest.TestCase):
    def check(self, text, expected, delimiter=',', name='LUNCH_PIN'):
        before = len(patterns()[name].findall(text))
        evidence, fs = scan(text, delimiter)
        self.assertEqual((before, *counts(evidence, name)), expected)
        result = [f for f in fs if f['pattern_name'] == name]
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0]['count'], expected[0]+expected[1]-expected[2])
        self.assertFalse(result[0]['header_only'])
        self.assertTrue(result[0]['column_bound'])
        self.assertFalse(result[0]['is_metadata'])
        self.assertEqual(len(patterns()[name].findall(text)), before)
        return fs

    def test_one_overlap_preserves_single_pin_warning(self):
        fs = self.check('lunch_pin\n0000\n', (1, 1, 1))
        self.assertEqual(fs[0]['confidence'], 'medium')
        self.assertEqual(e.worst_action(fs), 'warn')

    def test_two_rows_only_first_overlaps(self):
        self.check('lunch_pin\n0000\n0000\n', (1, 2, 1))

    def test_disjoint_equal_inline_value_is_not_deduplicated(self):
        self.check('lunch_pin,note\n0000,lunch_pin=0000\n', (1, 1, 0))

    def test_three_occurrences_catches_max_based_undercount(self):
        self.check('lunch_pin\n0000\n0000\nlunch_pin=0000\n', (2, 2, 1))

    def test_duplicate_columns_and_rows_count_physical_cells(self):
        self.check('lunch_pin,lunch_pin\n0000,0000\n0000,0000\n', (1, 4, 1))

    def test_tsv_label_in_previous_field_and_comma_control(self):
        for delimiter, expected in [('\t', (1, 1, 1)), (',', (0, 1, 0))]:
            self.check(delimiter.join(['note','lunch_pin'])+'\n'+delimiter.join(['lunch_pin','0000'])+'\n', expected, delimiter)

    def test_quotes_and_header_quotes_preserve_raw_regex(self):
        for text in ['lunch_pin\n"0000"\n', '"lunch_pin"\n0000\n', '"lunch_pin"\n"0000"\n']:
            self.check(text, (0, 1, 0))

    def test_bom_newline_whitespace_and_unicode_positions(self):
        digits = chr(0x660)*4
        for delimiter in [',', '\t']:
            for newline in ['\r','\n','\r\n']:
                text = '\ufeffnote'+delimiter+'lunch_pin'+newline+'"a\r\nb\r\"\"c"'+delimiter+'  '+digits+'  '+newline
                self.check(text, (0, 1, 0), delimiter)
            self.check('\ufefflunch_pin\r\n  '+digits+'  \r\n', (1, 1, 1), delimiter)

    def test_header_unknown_and_nonbinding_cells_never_subtract(self):
        for text in ['lunch_pin\n', 'note\nlunch_pin=0000\n', 'note,other\nlunch_pin,0000\n']:
            evidence, fs = scan(text)
            self.assertEqual(evidence, {})
            self.assertEqual([(f['pattern_name'],f['count']) for f in fs], [(f['pattern_name'],f['count']) for f in e.scan_content(text)])

    def test_terminal_decimal_assumption_for_all_three_patterns(self):
        for name, label, width in [('SASID','sasid',10),('STUDENT_ID_LABELED','student_id',4),('LUNCH_PIN','lunch_pin',4)]:
            for digit in ['0', chr(0x660)]:
                text = label+'='+digit*width
                match = patterns()[name].search(text)
                self.assertIsNotNone(match)
                self.assertEqual(c.numeric_tail(text, match, c.DIGIT_WIDTHS[name]), (len(label)+1,len(text)))
                self.check(label+'\n'+digit*width+'\n',(1,1,1),name=name)

    def test_partial_tail_and_different_pattern_do_not_subtract(self):
        text = 'lunch_pin=0000 sasid='+'0'*10
        cur = c.TailCursor(text, patterns()['LUNCH_PIN'], 6)
        start = text.index('=')+1
        self.assertEqual(cur.accept(start+1, start+4), 0)
        cur = c.TailCursor(text, patterns()['SASID'], 12)
        self.assertEqual(cur.accept(start,start+4), 0)

    def test_late_malformed_record_discards_both_aggregates(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'synthetic.csv';path.write_text('lunch_pin\n0000\n"bad')
            si=e.read_file_content(str(path))
        self.assertEqual(si.column_evidence,{})
        self.assertEqual(si.reader_error,'EXTRACTION_FAILED')
        self.assertTrue(e.scan_content(si.content))

    def test_skip_and_early_full_action_parity(self):
        text='student_id\n0000\n'
        evidence,_=scan(text)
        self.assertEqual(e.scan_content(text,column_evidence=evidence,skip_patterns={'STUDENT_ID_LABELED'}),[])
        self.assertEqual(e.worst_action(e.scan_content(text,column_evidence=evidence,early_exit=True)), e.worst_action(e.scan_content(text,column_evidence=evidence)))

    def test_invalid_aggregate_is_operational_hold_not_subtraction(self):
        for b,o in [(True,0),(1,True),(-1,0),(1,-1),(1,2),(2,2),(1.0,0)]:
            fs=e.scan_content('lunch_pin\n0000\n',column_evidence={'LUNCH_PIN':c.BoundEvidence(b,o)})
            self.assertEqual(e.worst_action(fs),'block')
            self.assertIn('SCAN_READER_UNAVAILABLE',{f['pattern_name'] for f in fs})
            self.assertEqual(next(f['count'] for f in fs if f['pattern_name']=='LUNCH_PIN'),1)

    def test_differential_generated_attribution_preserves_regex_counts(self):
        rng=random.Random(1213)
        for delimiter in [',','\t']:
            for _ in range(60):
                rows=[['note','lunch_pin']]+[[rng.choice(['lunch_pin','a\nb','"','']),rng.choice(['0000',' 0000 ','NULL'])] for _ in range(5)]
                buf=io.StringIO(newline='');csv.writer(buf,delimiter=delimiter).writerows(rows);text=buf.getvalue()
                evidence,fs=scan(text,delimiter)
                # Independent physical intervals from this test's numeric-only
                # fields: a legacy match ends on a qualifying value, never on NULL.
                regex_count=len(patterns()['LUNCH_PIN'].findall(text))
                bound=sum(row[1].strip().isdecimal() for row in rows[1:])
                overlap=regex_count
                if bound:
                    self.assertEqual(counts(evidence,'LUNCH_PIN'),(bound,overlap))
                    self.assertEqual(next(f['count'] for f in fs if f['pattern_name']=='LUNCH_PIN'),bound)
                else:self.assertEqual(evidence,{})


class AllocationTests(unittest.TestCase):
    def test_csv_trace_is_row_local_and_parser_admitted(self):
        real=c._parse_record;observed=[]
        def parsed(text,start,end,delimiter):
            observed.append((start,end));return real(text,start,end,delimiter)
        with mock.patch.object(c,'_parse_record',side_effect=parsed):
            rows=c.validated_rows('a,b\n0000,0000\n0000,0000\n',',',with_spans=True)
            previous_end=0
            for row,spans in rows:
                self.assertEqual(len(spans),len(row));self.assertEqual(len(spans),2)
                self.assertEqual(len(observed),previous_end+1)
                self.assertTrue(all(observed[-1][0]<=a<=b<=observed[-1][1] for a,b in spans))
                previous_end+=1
        self.assertEqual(previous_end,3)

    def test_csv_regex_cursors_remain_lazy(self):
        real=patterns()['LUNCH_PIN'];advances=[]
        class Pattern:
            def finditer(self,text):
                for match in real.finditer(text):
                    advances.append(1);yield match
        text='lunch_pin\n0000\n'+('lunch_pin=0000\n'*100)
        evidence=c.csv_evidence(text,',',{'LUNCH_PIN':Pattern()})
        self.assertEqual(counts(evidence,'LUNCH_PIN'),(1,1))
        self.assertLessEqual(len(advances),2)

    def test_compact_interval_cap_checked_before_append(self):
        store=c.IntervalStore(50_000)
        for i in range(50_000):store.add('LUNCH_PIN',2*i,2*i+1)
        self.assertEqual(store.payload_bytes,800_000)
        with self.assertRaises(c.ColumnarError):store.add('LUNCH_PIN',100_000,100_001)
        self.assertEqual(store.payload_bytes,800_000)
        self.assertTrue(all(isinstance(v,array) and v.itemsize==8 for v in store.intervals.values()))

    def test_bad_interval_and_offset_fail_without_append(self):
        for a,b in [(-1,2),(3,2),(0,1<<64),(True,2)]:
            store=c.IntervalStore(2)
            with self.assertRaises(c.ColumnarError):store.add('LUNCH_PIN',a,b)
            self.assertEqual(store.payload_bytes,0)


    def test_previous_csv_trace_and_row_released_before_next_allocation(self):
        class WeakList(list):
            pass
        prior = {'trace': None, 'row': None}
        def trace():
            self.assertTrue(prior['trace'] is None or prior['trace']() is None)
            result = WeakList()
            prior['trace'] = weakref.ref(result)
            return result
        original = c._parse_record
        def parse(*args):
            self.assertTrue(prior['row'] is None or prior['row']() is None)
            result = WeakList(original(*args))
            prior['row'] = weakref.ref(result)
            return result
        with mock.patch.object(c, '_field_spans', side_effect=trace), mock.patch.object(c, '_parse_record', side_effect=parse):
            result = c.csv_evidence('lunch_pin\n'+('0000\n'*100), ',', patterns())
        self.assertEqual(counts(result, 'LUNCH_PIN'), (100, 1))
        self.assertIsNone(prior['trace']())
        self.assertIsNone(prior['row']())

    def test_excess_columns_stop_at_first_capacity_check(self):
        visited = []
        class Source(str):
            def __getitem__(self, index):
                if isinstance(index, int): visited.append(index)
                return super().__getitem__(index)
        with mock.patch.object(c, '_parse_record') as parser, mock.patch.object(c, '_copy_record') as copy:
            with self.assertRaises(c.ColumnarError):
                c.csv_evidence(Source(','*100_000), ',', patterns())
            parser.assert_not_called(); copy.assert_not_called()
        self.assertEqual(max(visited), c.MAX_COLUMNS-1)


class SurfaceCacheTests(unittest.TestCase):
    def load(self, name, relative):
        spec = importlib.util.spec_from_file_location(name, ROOT/relative)
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
        return module

    def test_hook_mcp_report_and_cache_share_counts_without_overlap_payload(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); home = root/'home'; home.mkdir(); data = root/'data'; data.mkdir()
            path = data/'synthetic.csv'; path.write_text('lunch_pin\n0000\n')
            env = {k:v for k,v in os.environ.items() if not k.startswith(('FERPA_GUARD_', 'PII_GUARDIAN_'))}
            env.update(HOME=str(home), FERPA_GUARD_CACHE='1')
            payload = json.dumps({'tool_name':'Read','tool_input':{'file_path':str(path)}})
            for _ in range(2):
                result = subprocess.run([sys.executable,str(ROOT/'claude-code/pii_guardian.py')], input=payload, text=True, capture_output=True, env=env)
                self.assertEqual(result.returncode, 0)
                self.assertIn('Possible sensitive data', result.stderr)
            cache = json.loads((home/'.claude/ferpa-guard-cache.json').read_text())
            self.assertEqual(cache['version'], 8)
            stored = cache['entries'][0]['findings']
            self.assertEqual([(f['pattern_name'],f['count'],f['confidence']) for f in stored], [('LUNCH_PIN',1,'medium')])
            audit = (home/'.claude/logs/ferpa-guard-audit.jsonl').read_text()
            self.assertEqual([json.loads(line)['action'] for line in audit.splitlines()], ['warn','warn'])
            mcp = self.load('binding_mcp', 'cowork/mcp_server.py')
            report = self.load('binding_report', 'claude-code/pii_scan_report.py')
            with mock.patch.object(mcp,'AUDIT_LOG_PATH',home/'mcp-audit.log'), mock.patch.dict(os.environ,{'HOME':str(home)}):
                remote = mcp.scan_file(str(path))
                batch = report.scan_directory(str(data))
            self.assertEqual(remote['action'],'warn')
            public_keys = {'pattern_name','description','severity','confidence','count'}
            self.assertEqual(remote['findings'],[{k:v for k,v in f.items() if k in public_keys} for f in stored])
            self.assertEqual(batch['blocked'][0]['findings'],stored)
            # Legacy report grouping is preserved until the separate floor stage.
            self.assertEqual(e.worst_action(batch['blocked'][0]['findings']),'warn')
            serialized = json.dumps([cache, remote, batch])+audit+(home/'mcp-audit.log').read_text()
            for forbidden in ['regex_overlap_count','bound_count','intervals','pending']:
                self.assertNotIn(forbidden,serialized)

    def test_old_cache_poison_does_not_hide_bound_value(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);path=root/'synthetic.csv';path.write_text('student_number\n0000\n')
            stat=path.stat();cache_path=root/'.claude/ferpa-guard-cache.json';cache_path.parent.mkdir()
            env={k:v for k,v in os.environ.items() if not k.startswith(('FERPA_GUARD_','PII_GUARDIAN_'))}
            env.update(HOME=str(root),FERPA_GUARD_CACHE='1')
            for version in range(4,8):
                cache_path.write_text(json.dumps({'version':version,'entries':[{'key':[str(path),stat.st_mtime,stat.st_size],'cached_at':time.time(),'findings':[]}]}))
                result=subprocess.run([sys.executable,str(ROOT/'claude-code/pii_guardian.py')],input=json.dumps({'tool_name':'Read','tool_input':{'file_path':str(path)}}),text=True,capture_output=True,env=env)
                self.assertEqual(result.returncode,2)
                self.assertEqual(json.loads(cache_path.read_text())['version'],8)

    def test_attribution_uses_identical_content_and_compiled_patterns(self):
        original=c.csv_evidence;observed=[]
        def trace(content, delimiter, registry):
            observed.append(content)
            for name,pattern in registry.items():self.assertIs(pattern,e.PII_PATTERNS[name]['pattern'])
            return original(content,delimiter,registry)
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'synthetic.tsv';path.write_bytes(b'lunch_pin\r\n0000\r\n')
            with mock.patch.object(c,'csv_evidence',side_effect=trace):si=e.read_file_content(str(path))
        self.assertIs(observed[0],si.content)
        self.assertEqual(si.content,'lunch_pin\n0000\n')
        self.assertEqual(counts(si.column_evidence,'LUNCH_PIN'),(1,1))


class WorkbookOverlapTests(unittest.TestCase):
    def read(self, rows, other=None, property_text=None):
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'synthetic.xlsx';w=openpyxl.Workbook();w.properties.creator='SYNTH';w.active.title='SYNTH'
            for row in rows:w.active.append(row)
            if other:
                ws=w.create_sheet('OTHER')
                for row in other:ws.append(row)
            if property_text:w.properties.description=property_text
            w.save(path);w.close();return e.read_file_content(str(path))

    def test_exact_rendering_none_multiline_sheets_and_properties(self):
        si=self.read([['note',None,'lunch_pin'],[None,None,'0000'],['x\ny',None,'0000']], [['lunch_pin'],['0000']], 'lunch_pin=0000')
        expected='[Sheet: SYNTH]\nnote,lunch_pin\n0000\nx\ny,0000\n[Sheet: OTHER]\nlunch_pin\n0000\n[Property: description] lunch_pin=0000\n[Property: creator] SYNTH'
        self.assertEqual(si.content,expected)
        self.assertEqual(counts(si.column_evidence,'LUNCH_PIN'),(3,2))
        fs=e.scan_content(si.content,si.header_line_indices,column_evidence=si.column_evidence)
        self.assertEqual(next(f['count'] for f in fs if f['pattern_name']=='LUNCH_PIN'),4)
        self.assertNotIn('regex_overlap_count',json.dumps(fs))

    def test_numeric_renderings_and_empty_string_positions(self):
        for value in [4000,4000.0,'0000',' 0000 ']:
            si=self.read([['lunch_pin'],[value]])
            self.assertEqual(counts(si.column_evidence,'LUNCH_PIN'),(1,1))
        si=self.read([['note','lunch_pin'],['','0000']])
        self.assertEqual(counts(si.column_evidence,'LUNCH_PIN'),(1,1))

    def test_float_suffix_containment_uses_actual_rendering(self):
        text='lunch_pin\n4000.0'
        store=c.IntervalStore(1);store.add('LUNCH_PIN',text.index('\n')+1,len(text))
        evidence=store.reduce(text,patterns())
        self.assertEqual(counts(evidence,'LUNCH_PIN'),(1,1))
        self.assertEqual(store.payload_bytes,0)

    def test_comment_hold_discards_overlap(self):
        import openpyxl
        from openpyxl.comments import Comment
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'synthetic.xlsx';w=openpyxl.Workbook();w.active.append(['lunch_pin']);w.active.append(['0000']);w.active['A2'].comment=Comment('SYNTH','SYNTH');w.save(path);w.close()
            si=e.read_file_content(str(path))
        self.assertEqual(si.truncated,'XLSX_COMMENTS');self.assertEqual(si.column_evidence,{})

    def test_attribution_exception_retains_content_and_holds(self):
        with mock.patch.object(c.IntervalStore,'reduce',side_effect=c.ColumnarError):
            si=self.read([['lunch_pin'],['0000']])
        self.assertTrue(si.content);self.assertEqual(si.column_evidence,{})
        self.assertEqual(si.reader_error,'EXTRACTION_FAILED')


if __name__=='__main__':unittest.main()
