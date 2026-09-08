#!/usr/bin/env python3
"""Contract tests for XLSX comment hold and redactor verification (Phase 0, spec 2.3 and 2.4).

All fixtures are synthetic. Cell comments are invisible to the read-only
scanner, so the reader holds on any container that has, or might have,
comment parts, and the redactor refuses to keep an output it cannot prove
is comment-free.
"""

import contextlib
import io
import json
import os
import subprocess
import sys
import tempfile
import time
import types
import unittest
import warnings
import zipfile
from pathlib import Path
from unittest import mock

# Match tests/test_pii_guardian.py: add the project root for shared imports.
_project_root = Path(__file__).parent.parent
sys.path.insert(0, str(_project_root))

import shared.pii_engine as pii_engine
import shared.pii_redactor as redactor
from shared.pii_engine import (
    read_xlsx_file,
    scan_content,
    scan_incomplete_finding,
    worst_action,
    xlsx_comment_status,
)

try:
    import openpyxl
    from openpyxl.comments import Comment
except ImportError:  # pragma: no cover - the suite skips openpyxl fixtures
    openpyxl = None
    Comment = None


HOOK_SCRIPT = _project_root / "claude-code" / "pii_guardian.py"
SYNTHETIC_SSN = "123-45-6789"
XLSX_COMMENTS_TEXT = (
    "The spreadsheet has cell comments, or comments could not be ruled out, "
    "and the scanner cannot check them"
)
EXCEL_FORBIDDEN = set("\\/?*[]:")


def _requires_openpyxl(test):
    return unittest.skipIf(openpyxl is None, "openpyxl is required for workbook fixtures")(test)


def _write_workbook(path, rows, *, comment_at=None, comment_text=None,
                    titles=None, properties=None):
    """Write a synthetic workbook. `rows` is a list of row tuples for sheet 1.

    `titles` renames sheets (extra titles create extra sheets); `comment_at`
    attaches a legacy comment to that cell reference on sheet 1; `properties`
    sets core document properties.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(list(row))
    if titles:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            ws.title = titles[0]
            for extra in titles[1:]:
                wb.create_sheet(title=extra)
    if comment_at:
        ws[comment_at].comment = Comment(comment_text or "synthetic note", "Example Author")
    for name, value in (properties or {}).items():
        setattr(wb.properties, name, value)
    wb.save(path)
    wb.close()
    return path


def _write_zip(path, names):
    with zipfile.ZipFile(path, "w") as zf:
        for name in names:
            zf.writestr(name, "<synthetic/>")
    return path


def _run_hook(filepath, *, home, env_extra=None):
    env = os.environ.copy()
    env["HOME"] = str(home)
    if env_extra:
        env.update(env_extra)
    payload = json.dumps({"tool_name": "Read", "tool_input": {"file_path": str(filepath)}})
    return subprocess.run(
        [sys.executable, str(HOOK_SCRIPT)],
        input=payload,
        capture_output=True,
        text=True,
        env=env,
        timeout=30,
    )


class _BrokenZipModule(types.SimpleNamespace):
    """Stand-in for the zipfile module whose ZipFile always fails to open."""

    BadZipFile = zipfile.BadZipFile

    @staticmethod
    def ZipFile(*args, **kwargs):
        raise zipfile.BadZipFile("synthetic unreadable container")


class TestXlsxCommentStatus(unittest.TestCase):
    @_requires_openpyxl
    def test_openpyxl_commented_workbook_is_present(self):
        """Spec 2.3: the openpyxl layout xl/comments/comment1.xml is detected."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_workbook(
                Path(tmp) / "commented.xlsx", [("color",), ("blue",)],
                comment_at="A2", comment_text="synthetic note",
            )
            with zipfile.ZipFile(path) as zf:
                names = zf.namelist()
            self.assertIn("xl/comments/comment1.xml", names)
            self.assertEqual(xlsx_comment_status(str(path)), "present")

    def test_excel_legacy_layout_is_present(self):
        """Spec 2.3: the Excel layout xl/comments1.xml is detected."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_zip(
                Path(tmp) / "legacy.xlsx",
                ["[Content_Types].xml", "xl/workbook.xml", "xl/comments1.xml"],
            )
            self.assertEqual(xlsx_comment_status(str(path)), "present")

    def test_excel_threaded_layout_is_present(self):
        """Spec 2.3: xl/threadedComments/threadedComment1.xml is detected."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_zip(
                Path(tmp) / "threaded.xlsx",
                ["xl/workbook.xml", "xl/threadedComments/threadedComment1.xml"],
            )
            self.assertEqual(xlsx_comment_status(str(path)), "present")

    def test_case_insensitive_member_names(self):
        """Spec 2.3: the container check is case-insensitive."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_zip(Path(tmp) / "upper.xlsx", ["XL/Comments1.XML"])
            self.assertEqual(xlsx_comment_status(str(path)), "present")

    @_requires_openpyxl
    def test_clean_workbook_is_absent(self):
        """Spec 2.3: a workbook with no comment parts is absent."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_workbook(Path(tmp) / "clean.xlsx", [("color",), ("blue",)])
            self.assertEqual(xlsx_comment_status(str(path)), "absent")

    def test_vml_drawing_alone_is_absent(self):
        """The regex is anchored to comment parts, not to any name containing 'comments'."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_zip(
                Path(tmp) / "vml_only.xlsx",
                ["xl/workbook.xml", "xl/drawings/commentsDrawing1.vml", "xl/worksheets/sheet1.xml"],
            )
            self.assertEqual(xlsx_comment_status(str(path)), "absent")

    def test_non_zip_is_unknown(self):
        """Spec 2.3: BadZipFile yields unknown, never absent."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "corrupt.xlsx"
            path.write_bytes(b"this is deliberately not an OOXML archive")
            self.assertEqual(xlsx_comment_status(str(path)), "unknown")

    def test_nonexistent_path_is_unknown(self):
        """Spec 2.3: OSError yields unknown."""
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(xlsx_comment_status(str(Path(tmp) / "missing.xlsx")), "unknown")

    def test_directory_is_unknown(self):
        """Spec 2.3: any other open failure yields unknown."""
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(xlsx_comment_status(tmp), "unknown")

    def test_description_text_and_finding_mapping(self):
        """Spec 2.3: XLSX_COMMENTS is a known truncation code with the pinned wording."""
        self.assertEqual(pii_engine._TRUNCATION_DESCRIPTIONS["XLSX_COMMENTS"], XLSX_COMMENTS_TEXT)
        finding = scan_incomplete_finding("XLSX_COMMENTS", "/synthetic/records.xlsx")
        self.assertEqual(finding["pattern_name"], "SCAN_INCOMPLETE")
        self.assertEqual(finding["code"], "XLSX_COMMENTS")
        self.assertEqual(finding["description"], XLSX_COMMENTS_TEXT)
        self.assertEqual(worst_action([finding]), "block")


@_requires_openpyxl
class TestReaderCommentHold(unittest.TestCase):
    def test_commented_workbook_holds_and_keeps_cells(self):
        """Spec 2.3: comments set truncated=XLSX_COMMENTS; cells still scan; no extraction."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_workbook(
                Path(tmp) / "commented.xlsx", [("color", "count"), ("blue", 3)],
                comment_at="A2", comment_text=f"synthetic note {SYNTHETIC_SSN}",
            )
            result = read_xlsx_file(str(path))
        self.assertEqual(result.truncated, "XLSX_COMMENTS")
        self.assertEqual(result.reader_error, "")
        self.assertIn("blue,3", result.content)
        self.assertNotIn(SYNTHETIC_SSN, result.content)

    def test_ssn_only_in_comment_blocks_via_scan_incomplete(self):
        """Spec 2.3 and acceptance 6: a comment-only SSN yields a blocking hold, not a detection."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_workbook(
                Path(tmp) / "comment_ssn.xlsx", [("color", "count"), ("blue", 3)],
                comment_at="B2", comment_text=f"SSN {SYNTHETIC_SSN}",
            )
            result = read_xlsx_file(str(path))
            findings = scan_content(result.content, result.header_line_indices)
            if result.reader_error:
                findings.append(pii_engine.reader_error_finding(result.reader_error, str(path)))
            if result.truncated:
                findings.append(scan_incomplete_finding(result.truncated, str(path)))
        names = {f["pattern_name"] for f in findings}
        self.assertIn("SCAN_INCOMPLETE", names)
        self.assertNotIn("SSN", names)
        self.assertEqual(worst_action(findings), "block")
        codes = {f.get("code") for f in findings}
        self.assertIn("XLSX_COMMENTS", codes)

    def test_clean_workbook_has_no_hold(self):
        """A workbook without comment parts must not be held."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_workbook(Path(tmp) / "clean.xlsx", [("color", "count"), ("blue", 3)])
            result = read_xlsx_file(str(path))
        self.assertEqual(result.truncated, "")
        self.assertEqual(result.reader_error, "")

    def test_corrupt_container_reports_open_failed_and_hold(self):
        """Spec 2.3: unknown status holds even when openpyxl also fails; both findings surface."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "corrupt.xlsx"
            path.write_bytes(b"deliberately not an OOXML archive")
            result = read_xlsx_file(str(path))
        self.assertEqual(result.reader_error, "OPEN_FAILED")
        self.assertEqual(result.truncated, "XLSX_COMMENTS")
        self.assertEqual(result.content, "")

    def test_comment_hold_takes_precedence_over_cell_cap(self):
        """truncated is one code; the comment hold wins when both apply (spec 2.3 wording)."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_workbook(
                Path(tmp) / "capped.xlsx",
                [("color", "count"), ("blue", 3), ("red", 4)],
                comment_at="A2",
            )
            with mock.patch.object(pii_engine, "MAX_XLSX_CELLS", 1):
                result = read_xlsx_file(str(path))
        self.assertEqual(result.truncated, "XLSX_COMMENTS")

    def test_title_property_ssn_is_detected(self):
        """Spec 2.3: a synthetic SSN in the title property yields an SSN finding."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_workbook(
                Path(tmp) / "props.xlsx", [("color", "count"), ("blue", 3)],
                properties={"title": f"Roster {SYNTHETIC_SSN}"},
            )
            result = read_xlsx_file(str(path))
            findings = scan_content(result.content, result.header_line_indices)
        self.assertEqual(result.truncated, "")
        self.assertIn(f"[Property: title] Roster {SYNTHETIC_SSN}", result.content)
        names = {f["pattern_name"] for f in findings}
        self.assertIn("SSN", names)
        self.assertEqual(worst_action(findings), "block")

    def test_all_seven_properties_are_emitted_and_empty_ones_skipped(self):
        """Spec 2.3: the seven core properties appear as [Property: name] lines when non-empty."""
        values = {
            "title": "Example Title",
            "subject": "Example Subject",
            "description": "Example Description",
            "keywords": "example, keywords",
            "creator": "Example Creator",
            "lastModifiedBy": "Example Editor",
            "category": "Example Category",
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_workbook(
                Path(tmp) / "props.xlsx", [("color", "count"), ("blue", 3)], properties=values,
            )
            result = read_xlsx_file(str(path))
        for name, value in values.items():
            self.assertIn(f"[Property: {name}] {value}", result.content)
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_workbook(Path(tmp) / "bare.xlsx", [("color", "count"), ("blue", 3)])
            bare = read_xlsx_file(str(path))
        for name in ("title", "subject", "description", "keywords", "category"):
            self.assertNotIn(f"[Property: {name}]", bare.content)

    def test_property_lines_are_never_header_lines(self):
        """Spec 2.3: property lines are appended after the sheets and are not header rows."""
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_workbook(
                Path(tmp) / "props.xlsx", [("color", "count"), ("blue", 3)],
                properties={"title": "Example Title", "creator": "Example Creator"},
            )
            result = read_xlsx_file(str(path))
        lines = result.content.split("\n")
        self.assertTrue(result.header_line_indices, "the real header row should still be tagged")
        for idx in result.header_line_indices:
            self.assertFalse(lines[idx].startswith("[Property:"))
        property_indices = {i for i, line in enumerate(lines) if line.startswith("[Property:")}
        self.assertTrue(property_indices)
        self.assertTrue(property_indices.isdisjoint(result.header_line_indices))


@_requires_openpyxl
class TestRedactorXlsx(unittest.TestCase):
    def _redact(self, tmp, **kwargs):
        input_path = _write_workbook(Path(tmp) / "input.xlsx", kwargs.pop("rows"), **kwargs)
        output_path = Path(tmp) / "input_redacted.xlsx"
        count = redactor.redact_xlsx(input_path, output_path)
        return input_path, output_path, count

    def test_comments_are_stripped_and_output_is_absent(self):
        """Spec 2.4: the output has no comment parts and verifies absent."""
        with tempfile.TemporaryDirectory() as tmp:
            _, output_path, count = self._redact(
                tmp, rows=[("color", "count"), ("blue", 3)],
                comment_at="A2", comment_text=f"synthetic {SYNTHETIC_SSN}",
            )
            self.assertTrue(output_path.exists())
            self.assertEqual(xlsx_comment_status(str(output_path)), "absent")
            with zipfile.ZipFile(output_path) as zf:
                names = zf.namelist()
            self.assertFalse([n for n in names if pii_engine._XLSX_COMMENT_PARTS.match(n)])
            wb = openpyxl.load_workbook(output_path)
            self.assertIsNone(wb.active["A2"].comment)
            self.assertEqual(wb.active["A2"].value, "blue")
            self.assertEqual(wb.active["B2"].value, 3)
            wb.close()
        self.assertGreaterEqual(count, 1)

    def test_header_row_comment_is_stripped(self):
        """A comment on a header cell must not survive the header-row skip."""
        with tempfile.TemporaryDirectory() as tmp:
            _, output_path, _ = self._redact(
                tmp, rows=[("color", "count"), ("blue", 3)], comment_at="A1",
            )
            self.assertEqual(xlsx_comment_status(str(output_path)), "absent")

    def test_date_bearing_title_becomes_valid(self):
        """Spec 2.4: DOB 2012-01-01 redacts to a title with no '/' and at most 31 chars."""
        with tempfile.TemporaryDirectory() as tmp:
            _, output_path, _ = self._redact(
                tmp, rows=[("color", "count"), ("blue", 3)], titles=["DOB 2012-01-01"],
            )
            wb = openpyxl.load_workbook(output_path)
            titles = wb.sheetnames
            wb.close()
        self.assertEqual(len(titles), 1)
        title = titles[0]
        self.assertNotIn("2012-01-01", title)
        self.assertTrue(EXCEL_FORBIDDEN.isdisjoint(title), title)
        self.assertLessEqual(len(title), 31)

    def test_colliding_redacted_titles_are_distinct(self):
        """Spec 2.4: two sheets that redact to the same title end up distinct."""
        with tempfile.TemporaryDirectory() as tmp:
            _, output_path, _ = self._redact(
                tmp, rows=[("color", "count"), ("blue", 3)],
                titles=["DOB 2012-01-01", "DOB 2012-02-01"],
            )
            wb = openpyxl.load_workbook(output_path)
            titles = wb.sheetnames
            wb.close()
        self.assertEqual(len(titles), 2)
        self.assertEqual(len({t.lower() for t in titles}), 2)
        for title in titles:
            self.assertTrue(EXCEL_FORBIDDEN.isdisjoint(title), title)
            self.assertLessEqual(len(title), 31)
            self.assertNotIn("2012-0", title)

    def test_collision_suffix_stays_within_31_chars(self):
        """Spec 2.4: the de-duplication suffix fits inside the 31-character limit."""
        long_a = "DOB 2012-01-01 " + "x" * 19  # 34 chars; 31 after redaction
        long_b = "DOB 2012-03-01 " + "x" * 19
        with tempfile.TemporaryDirectory() as tmp:
            with warnings.catch_warnings():
                # openpyxl warns when it loads the deliberately over-long fixture titles.
                warnings.simplefilter("ignore")
                _, output_path, _ = self._redact(
                    tmp, rows=[("color", "count"), ("blue", 3)], titles=[long_a, long_b],
                )
            wb = openpyxl.load_workbook(output_path)
            titles = wb.sheetnames
            wb.close()
        self.assertEqual(len({t.lower() for t in titles}), 2)
        for title in titles:
            self.assertLessEqual(len(title), 31)
            self.assertTrue(EXCEL_FORBIDDEN.isdisjoint(title), title)

    def test_unchanged_title_keeps_its_name_when_another_collides(self):
        """A sheet whose title had no PII keeps it; the redacted sheet takes the suffix."""
        with tempfile.TemporaryDirectory() as tmp:
            _, output_path, _ = self._redact(
                tmp, rows=[("color", "count"), ("blue", 3)],
                titles=["DOB 2012-01-01", "DOB Q1-2012"],
            )
            wb = openpyxl.load_workbook(output_path)
            titles = wb.sheetnames
            wb.close()
        # Sheet-specific: the clean sheet keeps its exact name, the redacted
        # sheet takes the suffix (review finding 2026-09-07: the previous
        # assertions passed even if the two were swapped).
        self.assertEqual(titles[1], "DOB Q1-2012")
        self.assertEqual(titles[0], "DOB Q1-2012-2")

    def test_properties_are_redacted(self):
        """Spec 2.4: redact_text runs over the core properties."""
        with tempfile.TemporaryDirectory() as tmp:
            _, output_path, _ = self._redact(
                tmp, rows=[("color", "count"), ("blue", 3)],
                properties={
                    "title": f"Roster {SYNTHETIC_SSN}",
                    "description": "born 2012-01-01",
                    "creator": "Example Creator",
                },
            )
            wb = openpyxl.load_workbook(output_path)
            props = wb.properties
            wb.close()
        self.assertNotIn(SYNTHETIC_SSN, props.title)
        self.assertIn("XXX-XX-", props.title)
        self.assertNotIn("2012-01-01", props.description)
        self.assertEqual(props.creator, "Example Creator")

    def test_numeric_cells_are_unchanged(self):
        """Spec 2.4: numeric cells are left as they are, by design."""
        with tempfile.TemporaryDirectory() as tmp:
            _, output_path, _ = self._redact(
                tmp, rows=[("id", "ssn"), (123456789, SYNTHETIC_SSN)],
            )
            wb = openpyxl.load_workbook(output_path)
            ws = wb.active
            self.assertEqual(ws["A2"].value, 123456789)
            self.assertNotIn(SYNTHETIC_SSN, str(ws["B2"].value))
            wb.close()

    def test_verification_present_deletes_output_and_raises(self):
        """Spec 2.4: a present verdict after save refuses the output."""
        with tempfile.TemporaryDirectory() as tmp:
            input_path = _write_workbook(
                Path(tmp) / "input.xlsx", [("color", "count"), ("blue", 3)], comment_at="A2",
            )
            output_path = Path(tmp) / "input_redacted.xlsx"
            with mock.patch.object(redactor, "xlsx_comment_status", return_value="present"):
                with self.assertRaises(redactor.RedactionVerificationError) as ctx:
                    redactor.redact_xlsx(input_path, output_path)
            self.assertFalse(output_path.exists())
        self.assertIn("could not remove all comments", str(ctx.exception))

    def test_verification_unknown_deletes_output_and_raises(self):
        """Spec 2.4: an unreadable container after save (BadZipFile) refuses the output."""
        with tempfile.TemporaryDirectory() as tmp:
            input_path = _write_workbook(
                Path(tmp) / "input.xlsx", [("color", "count"), ("blue", 3)], comment_at="A2",
            )
            output_path = Path(tmp) / "input_redacted.xlsx"
            with mock.patch.object(pii_engine, "zipfile", _BrokenZipModule()):
                self.assertEqual(xlsx_comment_status(str(input_path)), "unknown")
                with self.assertRaises(redactor.RedactionVerificationError) as ctx:
                    redactor.redact_xlsx(input_path, output_path)
            self.assertFalse(output_path.exists())
        self.assertIn("could not be verified", str(ctx.exception))
        self.assertNotIn("could not remove", str(ctx.exception))

    def test_verification_failure_that_cannot_delete_names_the_leftover(self):
        """Review finding 2026-09-06 (GLM 5.2): when the unverified output cannot
        be deleted, the error must say so and name the file, never 'no output
        written'. The CLI and MCP tail follows output_remains."""
        with tempfile.TemporaryDirectory() as tmp:
            input_path = _write_workbook(
                Path(tmp) / "input.xlsx", [("color", "count"), ("blue", 3)], comment_at="A2",
            )
            output_path = Path(tmp) / "input_redacted.xlsx"
            with mock.patch.object(redactor, "xlsx_comment_status", return_value="present"), \
                    mock.patch.object(Path, "unlink", side_effect=OSError("locked")):
                with self.assertRaises(redactor.RedactionVerificationError) as ctx:
                    redactor.redact_xlsx(input_path, output_path)
            self.assertTrue(ctx.exception.output_remains)
            self.assertIn("remains at", str(ctx.exception))
            self.assertIn(str(output_path), str(ctx.exception))
            self.assertTrue(output_path.exists())
        # The normal path still reports deletion.
        err = redactor.RedactionVerificationError("Redaction could not remove all comments")
        self.assertFalse(err.output_remains)


@_requires_openpyxl
class TestRedactorCli(unittest.TestCase):
    def _run_main(self, argv):
        stdout = io.StringIO()
        with mock.patch.object(sys, "argv", argv), contextlib.redirect_stdout(stdout):
            try:
                redactor.main()
                code = 0
            except SystemExit as exc:
                code = exc.code
        return code, stdout.getvalue()

    def test_cli_success_prints_numeric_caveat(self):
        """Spec 2.4: the xlsx success line names what redaction does not change."""
        with tempfile.TemporaryDirectory() as tmp:
            input_path = _write_workbook(
                Path(tmp) / "input.xlsx", [("color", "count"), ("blue", 3)], comment_at="A2",
            )
            code, out = self._run_main(["pii_redactor.py", str(input_path)])
            self.assertEqual(code, 0)
            self.assertIn("Redacted", out)
            self.assertIn(
                "Numeric cells, names, and free text are not changed; see README, Redaction.", out,
            )
            self.assertTrue((Path(tmp) / "input_redacted.xlsx").exists())

    def test_cli_refusal_exits_nonzero_and_writes_nothing(self):
        """Spec 2.4: refusal prints the reason plus 'no output written.' and exits 1."""
        with tempfile.TemporaryDirectory() as tmp:
            input_path = _write_workbook(
                Path(tmp) / "input.xlsx", [("color", "count"), ("blue", 3)], comment_at="A2",
            )
            with mock.patch.object(redactor, "xlsx_comment_status", return_value="present"):
                code, out = self._run_main(["pii_redactor.py", str(input_path)])
            self.assertEqual(code, 1)
            self.assertIn("could not remove all comments", out)
            self.assertIn("no output written.", out)
            self.assertFalse((Path(tmp) / "input_redacted.xlsx").exists())

    def test_cli_refusal_on_unknown_status(self):
        """Spec 2.4 through the CLI: an uninspectable output is refused the same way."""
        with tempfile.TemporaryDirectory() as tmp:
            input_path = _write_workbook(
                Path(tmp) / "input.xlsx", [("color", "count"), ("blue", 3)], comment_at="A2",
            )
            with mock.patch.object(redactor, "xlsx_comment_status", return_value="unknown"):
                code, out = self._run_main(["pii_redactor.py", str(input_path)])
            self.assertEqual(code, 1)
            self.assertIn("could not be verified", out)
            self.assertIn("no output written.", out)
            self.assertFalse((Path(tmp) / "input_redacted.xlsx").exists())

    def test_cli_names_leftover_when_delete_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            input_path = _write_workbook(
                Path(tmp) / "input.xlsx", [("color", "count"), ("blue", 3)], comment_at="A2",
            )
            with mock.patch.object(redactor, "xlsx_comment_status", return_value="present"), \
                    mock.patch.object(Path, "unlink", side_effect=OSError("locked")):
                code, out = self._run_main(["pii_redactor.py", str(input_path)])
            self.assertEqual(code, 1)
            self.assertIn("remains at", out)
            self.assertNotIn("no output written", out)


@_requires_openpyxl
class TestMcpRedactFile(unittest.TestCase):
    def setUp(self):
        try:
            import cowork.mcp_server as server
        except ImportError:
            self.skipTest("mcp package is required for the MCP server tests")
        self.server = server
        self.tmpdir = tempfile.TemporaryDirectory()
        self.tmp = Path(self.tmpdir.name)
        self.audit_patch = mock.patch.object(
            server, "AUDIT_LOG_PATH", self.tmp / "audit" / "ferpa-guard-audit.log",
        )
        self.audit_patch.start()

    def tearDown(self):
        self.audit_patch.stop()
        self.tmpdir.cleanup()

    def test_redact_file_refusal_returns_error_without_output(self):
        """Spec 2.4: the MCP tool returns the refusal as an error field and leaves no file."""
        input_path = _write_workbook(
            self.tmp / "input.xlsx", [("color", "count"), ("blue", 3)], comment_at="A2",
        )
        with mock.patch.object(redactor, "xlsx_comment_status", return_value="present"):
            result = self.server.redact_file(str(input_path))
        self.assertIn("error", result)
        self.assertNotIn("output_path", result)
        self.assertIn("could not remove all comments", result["error"])
        self.assertFalse((self.tmp / "input_redacted.xlsx").exists())

    def test_redact_file_refusal_on_unknown_status(self):
        """Spec 2.4 through the MCP tool: an uninspectable output is refused."""
        input_path = _write_workbook(
            self.tmp / "input.xlsx", [("color", "count"), ("blue", 3)], comment_at="A2",
        )
        with mock.patch.object(redactor, "xlsx_comment_status", return_value="unknown"):
            result = self.server.redact_file(str(input_path))
        self.assertIn("error", result)
        self.assertNotIn("output_path", result)
        self.assertIn("could not be verified", result["error"])
        self.assertIn("no output written.", result["error"])
        self.assertFalse((self.tmp / "input_redacted.xlsx").exists())

    def test_redact_file_success_on_commented_workbook(self):
        """The MCP tool still produces a verified output when comments strip cleanly."""
        input_path = _write_workbook(
            self.tmp / "input.xlsx", [("color", "count"), ("blue", 3)], comment_at="A2",
        )
        result = self.server.redact_file(str(input_path))
        self.assertNotIn("error", result)
        self.assertEqual(xlsx_comment_status(result["output_path"]), "absent")


@_requires_openpyxl
class TestHookPoisonedCache(unittest.TestCase):
    def test_v3_clean_cache_for_commented_workbook_is_rejected(self):
        """A pre-comment-hold clean verdict is rejected at the current version."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            cache_dir = home / ".claude"
            cache_dir.mkdir(parents=True)
            path = _write_workbook(
                root / "commented_synthetic.xlsx", [("color", "count"), ("blue", 3)],
                comment_at="A2", comment_text=f"synthetic {SYNTHETIC_SSN}",
            )
            stat = path.stat()
            poisoned = {
                "version": 3,
                "entries": [{
                    "key": [str(path.resolve()), stat.st_mtime, stat.st_size],
                    "findings": [],
                    "cached_at": time.time(),
                }],
            }
            cache_file = cache_dir / "ferpa-guard-cache.json"
            cache_file.write_text(json.dumps(poisoned))
            result = _run_hook(path, home=home, env_extra={"FERPA_GUARD_CACHE": "1"})
            rewritten = json.loads(cache_file.read_text())

        self.assertEqual(result.returncode, 2)
        message = (result.stdout + result.stderr).lower()
        self.assertIn("could not check", message)
        self.assertIn("comments", message)
        self.assertNotIn(SYNTHETIC_SSN, result.stdout + result.stderr)
        self.assertEqual(rewritten["version"], 5)
        self.assertEqual(rewritten["entries"], [])

    def test_current_clean_cache_entry_is_honored(self):
        """Positive control for the test above (review finding 2026-09-07): an
        identical entry at the CURRENT version is loaded and short-circuits the
        scan, so the rejection above is provably the version check and not a
        broken or disabled cache."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            home = root / "home"
            cache_dir = home / ".claude"
            cache_dir.mkdir(parents=True)
            path = _write_workbook(
                root / "commented_synthetic.xlsx", [("color", "count"), ("blue", 3)],
                comment_at="A2", comment_text=f"synthetic {SYNTHETIC_SSN}",
            )
            stat = path.stat()
            current = {
                "version": 5,
                "entries": [{
                    "key": [str(path.resolve()), stat.st_mtime, stat.st_size],
                    "findings": [],
                    "cached_at": time.time(),
                }],
            }
            cache_file = cache_dir / "ferpa-guard-cache.json"
            cache_file.write_text(json.dumps(current))
            result = _run_hook(path, home=home, env_extra={"FERPA_GUARD_CACHE": "1"})

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertNotIn("comments", (result.stdout + result.stderr).lower())


if __name__ == "__main__":
    unittest.main()
