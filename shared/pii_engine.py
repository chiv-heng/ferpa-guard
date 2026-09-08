#!/usr/bin/env python3
"""
PII detection engine for K-12 student data protection.

Single source of truth for patterns, readers, scanning, confidence scoring,
and decision matrix. All surfaces (Claude Code hook, MCP server, redactor)
import from this module.
"""

import datetime
import json
import os
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from . import columnar


@dataclass
class ScanInput:
    """Content extracted from a file, with structured reader outcome state."""
    content: str
    header_line_indices: set = field(default_factory=set)
    reader_error: str = ""
    truncated: str = ""
    column_evidence: dict = field(default_factory=dict)

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

# Historical name retained for importer compatibility. This limit is measured
# in Unicode characters and remains the aggregate cap for DOCX extraction and
# the source of the per-page PDF character cap.
MAX_SCAN_BYTES = 2_000_000

# Max Unicode characters to scan per plain-text file. Text is read once as a
# bounded whole so scan_content still sees the complete available context.
MAX_SCAN_CHARS = 32_000_000

# Max non-empty cells to scan across an xlsx workbook
MAX_XLSX_CELLS = 50_000

# Max pages to scan per PDF (avoid stalling on bulk exports)
MAX_PDF_PAGES = 50

# File extensions to scan (data files likely to contain PII)
SCANNABLE_EXTENSIONS = {
    ".csv", ".tsv", ".txt", ".json", ".jsonl", ".xml",
    ".sql", ".log", ".dat",
    ".xls", ".xlsx",
    ".pdf", ".docx",
    ".md", ".html", ".htm",
}

# Directories to always skip (never contain student data)
SKIP_DIRS = {
    ".git", "node_modules", "__pycache__", ".next", ".vercel",
    ".claude", ".planning", ".backups",
}

# Documentation-convention basenames, exempt by exact (case-insensitive) name.
# Naming a column in docs is not disclosing a record; scanning these files
# blocked documentation ABOUT student data and locked files out of their own
# remediation. Exact-name match only -- never substring or glob.
SKIP_FILENAMES = {"claude.md", "readme.md", "agent.md", "agents.md"}

# Directory-walking allowfile: a dotfile dropped in a data directory (or any
# ancestor, up to the depth cap) declares allowed entries. Read fresh on every
# call, so allowing a file takes effect on the very next tool invocation --
# no settings edit, no session restart. This is the escape hatch environment
# variables can never be: the hook inherits Claude Code's environment, which
# a running session cannot change.
ALLOWFILE_NAME = ".pii-guardian-allow"
ALLOWFILE_MAX_DEPTH = 12

# ---------------------------------------------------------------------------
# PII Pattern Registry
# ---------------------------------------------------------------------------

PII_PATTERNS = {
    # --- Standard PII ---
    "SSN": {
        "pattern": re.compile(r"\b\d{3}-\d{2}-\d{4}\b"),
        "description": "Social Security Number (XXX-XX-XXXX)",
        "severity": "critical",
    },
    "SSN_NO_DASHES": {
        "pattern": re.compile(r"(?<!\d)\d{9}(?!\d)"),
        "description": "Possible SSN without dashes (9 consecutive digits)",
        "severity": "high",
        "min_context": True,
    },
    "EMAIL": {
        "pattern": re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b"),
        "description": "Email address",
        "severity": "medium",
    },
    "PHONE": {
        # Require delimiters between all three groups (dash, space, or parens).
        # Dots are excluded as delimiters to avoid false positives on decimal
        # numbers like 619066.6169 and 430163.2581 in financial data.
        # Matches: (401) 555-1234, 401-555-1234, 401 555 1234, +1 401-555-1234
        # Rejects: 6190666169, 619066.6169, 430163.2581, 401.555.1234
        "pattern": re.compile(
            r"(?<!\d)"                             # no digit before
            r"(?:\+?1[-\s])?"                      # optional country code
            r"(?:"
            r"\(\d{3}\)[-\s]?\d{3}[-\s]?\d{4}"    # (401) 555-1234
            r"|"
            r"\d{3}[-\s]\d{3}[-\s]\d{4}"          # 401-555-1234 or 401 555 1234
            r")"
            r"(?!\d)"                              # no digit after
        ),
        "description": "Phone number",
        "severity": "medium",
    },
    "DOB": {
        # Match MM/DD/YYYY, MM-DD-YYYY, and YYYY-MM-DD (ISO 8601) when near
        # a birth-related keyword. ISO format is common in SIS API exports
        # (PowerSchool, Infinite Campus). This avoids flagging random dates.
        "pattern": re.compile(
            r"\b(?:"
            r"(?:0[1-9]|1[0-2])[/-](?:0[1-9]|[12]\d|3[01])[/-](?:19|20)\d{2}"  # MM/DD/YYYY
            r"|"
            r"(?:19|20)\d{2}-(?:0[1-9]|1[0-2])-(?:0[1-9]|[12]\d|3[01])"        # YYYY-MM-DD
            r")\b"
        ),
        "description": "Date of birth",
        "severity": "high",
        "min_context": True,
        "context_keywords": {"birth", "dob", "born", "date_of_birth", "date of birth", "birthdate", "bday"},
    },

    # --- Student Identifiers ---
    "SASID": {
        "pattern": re.compile(r"\b(?:SASID|sasid)[\s:=]*\d{6,12}\b"),
        "description": "State-Assigned Student ID (SASID)",
        "severity": "critical",
    },
    "STUDENT_ID_LABELED": {
        "pattern": re.compile(r"(?i)\b(?:student[_\s]?id|pupil[_\s]?id|sis[_\s]?id|ps[_\s]?id|dcid)[\s:=]*\d{4,10}\b"),
        "description": "Labeled student identifier",
        "severity": "critical",
    },
    "LUNCH_PIN": {
        "pattern": re.compile(r"(?i)\b(?:lunch[_\s]?pin|meal[_\s]?pin|cafeteria[_\s]?pin)[\s:=]*\d{4,6}\b"),
        "description": "Lunch/meal PIN",
        "severity": "high",
    },

    # --- FERPA-Protected Fields ---
    "IEP_504_FLAG": {
        "pattern": re.compile(r"(?i)(?<![^\W_])(?:iep|504[_\s]?plan|504|sped|individualized[_\s]?education|accommodation[_\s]?plan)(?![^\W_])"),
        "description": "IEP/504 plan reference (FERPA-protected)",
        "severity": "high",
        "min_context": True,
    },
    "DISCIPLINE_RECORD": {
        "pattern": re.compile(r"(?i)(?<![^\W_])(?:suspen(?:sion|ded)|expel(?:led|sion)|disciplin(?:e|ary)[_\s]?(?:record|action|incident)|in[_\s]?school[_\s]?suspension|iss|oss)(?![^\W_])"),
        "description": "Disciplinary record reference",
        "severity": "high",
        "min_context": True,
    },
    "MEDICAL_INFO": {
        "pattern": re.compile(r"(?i)(?<![^\W_])(?:diagnos(?:is|ed)|medication|allergy|anaphyla|epinephrine|inhaler|seizure|diabetes|insulin)(?![^\W_])"),
        "description": "Medical information",
        "severity": "high",
        "min_context": True,
    },
    "PARENT_GUARDIAN": {
        "pattern": re.compile(r"(?i)\b(?:parent[_\s]?(?:name|email|phone|address|contact)|guardian[_\s]?(?:name|email|phone|address|contact)|mother[_\s]?(?:name|email)|father[_\s]?(?:name|email)|emergency[_\s]?contact)\b"),
        "description": "Parent/guardian contact information field",
        "severity": "high",
    },
    "HOME_ADDRESS": {
        # Case-insensitive to catch ALL CAPS and lowercase SIS exports.
        # Expanded road types for better coverage of US address formats.
        "pattern": re.compile(r"(?i)\b\d{1,5}\s+(?:[a-z]{2,}\s+){1,3}(?:st(?:reet)?|ave(?:nue)?|blvd|boulevard|dr(?:ive)?|ln|lane|rd|road|ct|court|way|pl(?:ace)?|cir(?:cle)?|ter(?:race)?|pkwy|parkway|hwy|highway|tr(?:ai)?l|loop|run|pike|al(?:le)?y)\b"),
        "description": "Home/street address",
        "severity": "medium",
    },
}

# Default context keywords for min_context patterns (education context)
DEFAULT_CONTEXT_KEYWORDS = {
    "student", "pupil", "child", "minor", "enrollment",
    "grade", "school", "classroom", "teacher", "parent",
    "guardian", "record", "ferpa", "confidential",
}

# ---------------------------------------------------------------------------
# File readers
# ---------------------------------------------------------------------------

def read_text_file(filepath: str) -> ScanInput:
    """Read bounded plain text and report any verified unread remainder."""
    try:
        f = open(filepath, "r", errors="replace")
    except Exception:
        return ScanInput(content="", reader_error="OPEN_FAILED")

    content = ""
    try:
        with f as handle:
            content = handle.read(MAX_SCAN_CHARS)
            truncated = "TEXT_LIMIT" if handle.read(1) else ""
        return ScanInput(content=content, truncated=truncated)
    except Exception:
        return ScanInput(content=content, reader_error="EXTRACTION_FAILED")


def _looks_like_header(row) -> bool:
    """Heuristic: a row looks like a header if all non-empty cells are
    short text strings (no numbers, no PII-like patterns).

    This prevents data rows from being falsely tagged as headers when
    the xlsx has no header row or headers start below row 1.
    """
    non_empty = [cell for cell in row if cell is not None]
    if not non_empty:
        return False

    for cell in non_empty:
        # Numeric cells are data, not headers
        if isinstance(cell, (int, float)):
            return False
        s = str(cell)
        # Very long values are data, not column labels
        if len(s) > 60:
            return False
        # If a cell matches a PII pattern (SSN, email, phone), it's data
        if re.search(r"\b\d{3}-\d{2}-\d{4}\b", s):
            return False
        if re.search(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b", s):
            return False
        if re.search(r"\b\d{3}[-\s]\d{3}[-\s]\d{4}\b", s):
            return False

    return True


# OOXML parts that carry cell comments. Excel writes legacy comments as
# xl/comments1.xml and threaded comments under xl/threadedComments/; openpyxl
# writes xl/comments/comment1.xml. The read-only scanner never loads any of
# them, so their presence is verified omission (spec 2.3).
_XLSX_COMMENT_PARTS = re.compile(
    r"^xl/(comments\d*\.xml|comments/[^/]+\.xml|threadedComments/[^/]+\.xml)$", re.I
)

# Core document properties scanned as [Property: <name>] lines (spec 2.3).
# The redactor imports this so both surfaces cover the same seven fields.
XLSX_CORE_PROPERTIES = (
    "title", "subject", "description", "keywords", "creator", "lastModifiedBy", "category",
)


def xlsx_comment_status(filepath) -> str:
    """Inspect the xlsx container for comment parts without loading the workbook.

    Returns "present" when any member is a comment part, "absent" when the
    member list was read completely and none matched, and "unknown" when the
    container could not be opened or listed (BadZipFile, OSError, anything
    else). Callers must treat "unknown" like "present": comments cannot be
    ruled out.
    """
    try:
        with zipfile.ZipFile(filepath) as zf:
            names = zf.namelist()
    except Exception:
        return "unknown"
    for name in names:
        if _XLSX_COMMENT_PARTS.match(name):
            return "present"
    return "absent"


def _xlsx_property_lines(wb) -> list[str]:
    """Render the non-empty core properties of an openpyxl workbook."""
    props = wb.properties
    lines = []
    for name in XLSX_CORE_PROPERTIES:
        value = getattr(props, name, None)
        if value is None:
            continue
        text = str(value).strip()
        if text:
            lines.append(f"[Property: {name}] {text}")
    return lines


def _binding_patterns():
    return {name: PII_PATTERNS[name]["pattern"] for name in columnar.DIGIT_WIDTHS}


def read_xlsx_file(filepath: str) -> ScanInput:
    """Extract cell text from an xlsx file using openpyxl.

    Returns a ScanInput with all cell values concatenated with newlines,
    preserving sheet names as context. The first row of each sheet is
    tagged as a header row only if it passes a heuristic check (all short
    text strings, no numeric or PII-like values). Scans up to
    MAX_XLSX_CELLS total, then appends [Property: <name>] lines for the
    non-empty core document properties (never header lines).

    Cell comments are never loaded in read-only mode. When the container
    holds comment parts, or cannot be inspected, the result is held with
    truncated="XLSX_COMMENTS" (spec 2.3). That hold takes precedence over
    XLSX_CELLS because `truncated` carries one code; either code blocks.
    """
    try:
        import openpyxl
    except ImportError:
        return ScanInput(content="", reader_error="MISSING_DEPENDENCY:openpyxl")

    # Container check runs before any workbook load so the hold is recorded
    # even when openpyxl cannot open the file; both findings then surface.
    comment_hold = "XLSX_COMMENTS" if xlsx_comment_status(filepath) != "absent" else ""

    try:
        identity = os.stat(filepath)
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    except Exception:
        return ScanInput(content="", reader_error="OPEN_FAILED", truncated=comment_hold)

    lines = []
    cell_count = 0
    header_line_indices = set()
    reader_error = ""
    truncated = ""
    evidence = {}
    intervals = columnar.IntervalStore(MAX_XLSX_CELLS)
    budget = {}
    character_cursor = 0

    def append_line(line):
        nonlocal character_cursor
        lines.append(line)
        character_cursor += len(line) + 1

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            append_line(f"[Sheet: {sheet_name}]")
            max_row, max_column = columnar.xlsx_bounds(wb, ws, budget)
            mapping = {}
            if not max_row or not max_column:
                continue
            for row_idx, row in enumerate(ws.iter_rows(min_row=1, min_col=1,
                    max_row=max_row, max_col=max_column, values_only=True)):
                row_vals = []
                row_length = 0
                for cell_index, cell in enumerate(row):
                    if cell is not None:
                        if cell_count >= MAX_XLSX_CELLS:
                            truncated = "XLSX_CELLS"
                            break
                        rendered = str(cell)
                        if row_vals:
                            row_length += 1  # The actual flattened comma.
                        start = character_cursor + row_length
                        row_vals.append(rendered)
                        row_length += len(rendered)
                        cell_count += 1
                        pattern = mapping.get(cell_index) if row_idx else None
                        if pattern and columnar.qualifying(pattern, cell, rendered=rendered):
                            intervals.add(pattern, start, start + len(rendered))
                if row_vals:
                    if row_idx == 0 and _looks_like_header(row):
                        header_line_indices.add(len(lines))
                        mapping = columnar.bindings(row)
                    append_line(",".join(row_vals))
                if truncated:
                    break
            if truncated:
                break
        for line in _xlsx_property_lines(wb):
            append_line(line)
        after = os.stat(filepath)
        if (identity.st_dev, identity.st_ino, identity.st_size, identity.st_mtime_ns) != (
                after.st_dev, after.st_ino, after.st_size, after.st_mtime_ns):
            reader_error = "EXTRACTION_FAILED"
    except Exception:
        reader_error = "EXTRACTION_FAILED"
    finally:
        try:
            wb.close()
        except Exception:
            if not reader_error:
                reader_error = "EXTRACTION_FAILED"

    if comment_hold:
        truncated = comment_hold
    content = "\n".join(lines)
    try:
        if not reader_error and not truncated:
            evidence = intervals.reduce(content, _binding_patterns())
    except Exception:
        reader_error = reader_error or "EXTRACTION_FAILED"
        evidence = {}
    finally:
        intervals.clear()

    return ScanInput(
        content=content,
        header_line_indices=header_line_indices,
        reader_error=reader_error,
        truncated=truncated,
        column_evidence={} if reader_error or truncated else evidence,
    )


def read_pdf_file(filepath: str) -> ScanInput:
    """Extract text from a PDF using PyMuPDF (fitz).

    Scans up to MAX_PDF_PAGES. Returns page text concatenated with
    page markers for context.
    """
    try:
        import fitz
    except ImportError:
        return ScanInput(content="", reader_error="MISSING_DEPENDENCY:pymupdf")

    try:
        doc = fitz.open(filepath)
    except Exception:
        return ScanInput(content="", reader_error="OPEN_FAILED")

    lines = []
    reader_error = ""
    truncated = ""
    try:
        if doc.needs_pass or doc.is_encrypted:
            return ScanInput(content="", reader_error="ENCRYPTED")

        page_count = doc.page_count
        if page_count > MAX_PDF_PAGES:
            truncated = "PDF_PAGES"

        page_char_limit = MAX_SCAN_BYTES // MAX_PDF_PAGES
        for i in range(min(page_count, MAX_PDF_PAGES)):
            page = doc[i]
            text = page.get_text()
            if text.strip():
                lines.append(f"[Page {i + 1}]")
                lines.append(text[:page_char_limit])
                if len(text) > page_char_limit and not truncated:
                    truncated = "PDF_PAGE_LIMIT"
    except Exception:
        reader_error = "EXTRACTION_FAILED"
    finally:
        try:
            doc.close()
        except Exception:
            if not reader_error:
                reader_error = "EXTRACTION_FAILED"

    return ScanInput(
        content="\n".join(lines),
        reader_error=reader_error,
        truncated=truncated,
    )


def read_docx_file(filepath: str) -> ScanInput:
    """Extract text from a .docx file using python-docx.

    Reads paragraph text and table cell text, capped at MAX_SCAN_BYTES.
    """
    try:
        from docx import Document
    except ImportError:
        return ScanInput(content="", reader_error="MISSING_DEPENDENCY:python-docx")

    try:
        doc = Document(filepath)
    except Exception:
        return ScanInput(content="", reader_error="OPEN_FAILED")

    parts = []
    total = 0

    reader_error = ""
    truncated = ""

    def add_text(text):
        nonlocal total, truncated
        if not text.strip():
            return
        if total >= MAX_SCAN_BYTES:
            truncated = "DOCX_LIMIT"
            return
        parts.append(text)
        total += len(text)

    try:
        for para in doc.paragraphs:
            add_text(para.text)
            if truncated:
                break

        if not truncated:
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text for cell in row.cells if cell.text.strip()]
                    if cells:
                        add_text(",".join(cells))
                    if truncated:
                        break
                if truncated:
                    break
    except Exception:
        reader_error = "EXTRACTION_FAILED"

    return ScanInput(
        content="\n".join(parts),
        reader_error=reader_error,
        truncated=truncated,
    )


def read_file_content(filepath: str) -> ScanInput:
    """Read file content using the appropriate reader for the file type."""
    ext = Path(filepath).resolve().suffix.lower()

    if ext == ".xlsx":
        return read_xlsx_file(filepath)
    elif ext == ".pdf":
        return read_pdf_file(filepath)
    elif ext == ".docx":
        return read_docx_file(filepath)
    elif ext == ".xls":
        # .xls (legacy format) not supported by openpyxl; scan as binary text
        return read_text_file(filepath)
    else:
        result = read_text_file(filepath)
        if ext in (".csv", ".tsv"):
            try:
                evidence = columnar.csv_evidence(result.content, ',' if ext == '.csv' else '\t', _binding_patterns())
                if not result.reader_error and not result.truncated:
                    result.column_evidence = evidence
            except Exception:
                result.column_evidence = {}
                result.reader_error = result.reader_error or "EXTRACTION_FAILED"
        return result


# ---------------------------------------------------------------------------
# Scanning
# ---------------------------------------------------------------------------

def write_audit_event(log_path, action: str, file: str, patterns, source: str = None, *,
                      floor: str = "high", strict: bool = False, policy_error: str = "") -> None:
    """Append one JSONL audit record. Never raises; never affects the decision.

    Records pattern NAMES only -- matched values and file content must never
    reach the audit trail. `source` is set on bypass events to attribute the
    allowlist mechanism that granted it (env / user file / ancestor allowfile).
    """
    rec = {
        "v": 1,
        "ts": datetime.datetime.now().isoformat(timespec="seconds"),
        "action": action,
        "file": file,
        "patterns": list(patterns),
        "floor": floor,
        "strict": strict,
    }
    if policy_error:
        rec["policy_error"] = "POLICY_CONFIG_INVALID"
    if source:
        rec["source"] = source
    try:
        log_path = Path(log_path)
        log_path.parent.mkdir(parents=True, exist_ok=True)
        with open(log_path, "a") as fh:
            fh.write(json.dumps(rec) + "\n")
    except OSError:
        pass


def collect_allowfile_entries(filepath: str) -> dict:
    """Collect allowed entries from `.pii-guardian-allow` files in the target's
    ancestor directories.

    Returns {resolved_entry_path: declaring_allowfile_path}. Attribution is
    nearest-first-wins: when nested allowfiles declare the same resolved entry,
    the entry maps to the allowfile closest to the target (setdefault), so
    audit provenance always names the nearest declaration.

    Entries are one per line; blank lines and `#` comments are ignored.
    Absolute entries are kept as written; relative entries resolve against the
    declaring allowfile's own directory. The target is resolved first, so the
    walk runs over the real file's ancestors even when read through a symlink.
    """
    entries: dict = {}
    try:
        target = Path(filepath).resolve()
    except OSError:
        return entries

    for parent in list(target.parents)[:ALLOWFILE_MAX_DEPTH]:
        allowfile = parent / ALLOWFILE_NAME
        try:
            if not allowfile.is_file():
                continue
            for line in allowfile.read_text(errors="replace").splitlines():
                line = line.strip()
                if not line or line.startswith("#"):
                    continue
                candidate = Path(line) if os.path.isabs(line) else (parent / line)
                try:
                    resolved_entry = str(candidate.resolve())
                except OSError:
                    resolved_entry = str(candidate)
                entries.setdefault(resolved_entry, str(allowfile))
        except OSError:
            continue

    return entries


def should_scan(filepath: str) -> bool:
    """Decide whether a file path is worth scanning."""
    p = Path(filepath)

    if p.suffix.lower() not in SCANNABLE_EXTENSIONS:
        return False

    if p.name.lower() in SKIP_FILENAMES:
        return False

    for part in p.parts:
        if part in SKIP_DIRS:
            return False

    return True


def scan_content(content: str, header_line_indices: set = None,
                  early_exit: bool = False,
                  skip_patterns: set = None, *,
                  column_evidence: dict = None) -> list[dict]:
    """Scan text content for PII patterns. Returns list of findings.

    Each finding includes confidence and header_only fields. When
    header_line_indices is provided (from XLSX reader), matches are
    tracked per-line to determine if they appear only in header rows.

    When early_exit=True, returns immediately after the first CRITICAL
    finding (used by the hook for faster blocking on large files).

    When skip_patterns is provided, those pattern names are excluded
    from scanning (e.g., {"SSN", "SSN_NO_DASHES"}).
    """
    if not content and not column_evidence:
        return []

    evidence = column_evidence or {}
    binding_regex_counts = {}
    evidence_error = False
    try:
        if not isinstance(evidence, dict) or len(evidence) > 3:
            raise columnar.ColumnarError()
        for name, item in evidence.items():
            if (name not in columnar.DIGIT_WIDTHS or not isinstance(item, columnar.BoundEvidence)
                    or type(item.bound_count) is not int
                    or type(item.regex_overlap_count) is not int
                    or not 0 <= item.regex_overlap_count <= item.bound_count):
                raise columnar.ColumnarError()
            # This is the scanner's ordinary regex count for this pattern,
            # computed once before early exit so invalid evidence cannot hide.
            count = sum(1 for _ in PII_PATTERNS[name]["pattern"].finditer(content))
            binding_regex_counts[name] = count
            if item.regex_overlap_count > count:
                raise columnar.ColumnarError()
    except Exception:
        evidence = {}
        evidence_error = True
    # Callers must pair evidence with the exact reader content and registry.
    # On invalid counts, preserve all regex detections and append a safe hold.
    findings = []
    content_lower = content.lower()

    # Check for default education context (used by most min_context patterns)
    has_edu_context = any(kw in content_lower for kw in DEFAULT_CONTEXT_KEYWORDS)

    # Split into lines for per-line matching when header info is available
    lines = content.split("\n") if header_line_indices else None

    for name, spec in PII_PATTERNS.items():
        if skip_patterns and name in skip_patterns:
            continue
        # Handle context requirements
        item = evidence.get(name)
        bound_count = item.bound_count if item else 0
        overlap_count = item.regex_overlap_count if item else 0
        if spec.get("min_context") and not bound_count:
            # Pattern-specific context keywords override the default set
            pattern_keywords = spec.get("context_keywords")
            if pattern_keywords:
                # This pattern needs its OWN keywords present (e.g., DOB needs "birth"/"dob")
                if not any(kw in content_lower for kw in pattern_keywords):
                    continue
            else:
                # Use default education context
                if not has_edu_context:
                    continue

        compiled = spec["pattern"]
        regex_count = (binding_regex_counts[name] if name in binding_regex_counts
                       else len(compiled.findall(content)))
        if regex_count or bound_count:
            # Determine if all matches are in header lines only
            header_only = False
            if header_line_indices and lines and not bound_count:
                has_data_match = False
                has_header_match = False
                for line_idx, line in enumerate(lines):
                    if compiled.search(line):
                        if line_idx in header_line_indices:
                            has_header_match = True
                        else:
                            has_data_match = True
                header_only = has_header_match and not has_data_match

            findings.append({
                "pattern_name": name,
                "description": spec["description"],
                "severity": spec["severity"],
                "count": regex_count + bound_count - overlap_count,
                "confidence": "high",
                "header_only": header_only,
                "is_metadata": name in METADATA_PATTERNS,
                "column_bound": bool(bound_count),
            })

            # Early exit: skip remaining patterns once a CRITICAL match is found
            if early_exit and spec["severity"] == "critical" and not evidence_error:
                return findings

    # Calculate confidence scores based on count, co-occurrence, and position
    if findings:
        total_lines = content.count("\n") + 1 if content else 0
        _calculate_confidence(findings, total_lines=total_lines)

    if evidence_error:
        findings.append(reader_error_finding("EXTRACTION_FAILED", ""))
    return findings


_READER_ERROR_DESCRIPTIONS = {
    "MISSING_DEPENDENCY:openpyxl": "The spreadsheet scanner (openpyxl) is not installed",
    "MISSING_DEPENDENCY:pymupdf": "The PDF scanner (PyMuPDF) is not installed",
    "MISSING_DEPENDENCY:python-docx": "The document scanner (python-docx) is not installed",
    "OPEN_FAILED": "The file could not be opened safely",
    "ENCRYPTED": "The PDF is encrypted or password-protected",
    "EXTRACTION_FAILED": "The scanner stopped while extracting the file",
    "UNREADABLE": "The file is unreadable",
}

_TRUNCATION_DESCRIPTIONS = {
    "TEXT_LIMIT": "The text scan limit left part of the file unchecked",
    "XLSX_CELLS": "The spreadsheet cell limit left part of the file unchecked",
    "PDF_PAGES": "The PDF page limit left part of the file unchecked",
    "PDF_PAGE_LIMIT": "The PDF per-page text limit left part of the file unchecked",
    "DOCX_LIMIT": "The document text limit left part of the file unchecked",
    "XLSX_COMMENTS": "The spreadsheet has cell comments, or comments could not be ruled out, and the scanner cannot check them",
}


def reader_error_finding(code: str, filepath: str) -> dict:
    """Create a sanitized blocking finding for a reader failure."""
    safe_code = code if code in _READER_ERROR_DESCRIPTIONS else "UNREADABLE"
    return {
        "pattern_name": "SCAN_READER_UNAVAILABLE",
        "description": _READER_ERROR_DESCRIPTIONS[safe_code],
        "severity": "high",
        "confidence": "high",
        "count": 1,
        "header_only": False,
        "is_metadata": False,
        "code": safe_code,
    }


def scan_incomplete_finding(code: str, filepath: str) -> dict:
    """Create a sanitized blocking finding for verified omitted content."""
    safe_code = code if code in _TRUNCATION_DESCRIPTIONS else "TEXT_LIMIT"
    return {
        "pattern_name": "SCAN_INCOMPLETE",
        "description": _TRUNCATION_DESCRIPTIONS[safe_code],
        "severity": "high",
        "confidence": "high",
        "count": 1,
        "header_only": False,
        "is_metadata": False,
        "code": safe_code,
    }


# Patterns whose matches are structurally unambiguous (always HIGH confidence)
_HIGH_CONFIDENCE_FLOOR = {"SSN", "SASID", "STUDENT_ID_LABELED"}

# Keyword-only patterns that match common words in non-PII contexts
# (compliance docs, training materials, policy documents, health class assignments).
# These require stronger co-occurrence signals to reach HIGH confidence because
# the keywords alone don't prove student-specific PII is present.
_KEYWORD_AMBIGUITY_PENALTY = {"IEP_504_FLAG", "DISCIPLINE_RECORD", "MEDICAL_INFO"}


def _calculate_confidence(findings: list[dict], total_lines: int = 0) -> None:
    """Compute confidence tier for each finding based on count,
    co-occurrence with other pattern types, header position, and
    pattern reliability.
    Mutates findings in place.
    """
    distinct_types = len(findings)

    for f in findings:
        # Hard floor: structurally unambiguous patterns are always HIGH
        if f["pattern_name"] in _HIGH_CONFIDENCE_FLOOR:
            f["confidence"] = "high"
            continue

        score = 0

        # Signal 1: match count
        if f["count"] >= 5:
            score += 3
        elif f["count"] >= 3:
            score += 2
        elif f["count"] >= 2:
            score += 1

        # Signal 2: co-occurrence with other PII types
        if distinct_types >= 3:
            score += 2
        elif distinct_types >= 2:
            score += 1

        # Signal 3: header position penalty
        if f["header_only"]:
            score -= 3

        # Signal 4: keyword-ambiguity penalty for patterns that match
        # common words (IEP in policy docs, "medication" in health class).
        # These patterns need co-occurrence with structural PII (SSN,
        # student ID, email) to be confident, not just high count alone.
        if f["pattern_name"] in _KEYWORD_AMBIGUITY_PENALTY:
            score -= 1

        # Signal 5: small-file boost
        # In files with fewer than 10 rows, even a single match is
        # proportionally significant. Boost to prevent underweighting.
        if 0 < total_lines < 10:
            score += 2

        # Map score to tier
        if score >= 3:
            f["confidence"] = "high"
        elif score >= 1:
            f["confidence"] = "medium"
        else:
            f["confidence"] = "low"


# Topic/field-name patterns: their regexes match words ABOUT student data
# (column names, policy vocabulary), not data values. A file containing only
# these is documentation, not disclosure. Unknown patterns default to VALUE,
# so a newly added pattern fails closed (blocking) rather than open.
METADATA_PATTERNS = {
    "IEP_504_FLAG",
    "DISCIPLINE_RECORD",
    "MEDICAL_INFO",
    "PARENT_GUARDIAN",
}


def _is_metadata(pattern_name: str) -> bool:
    return pattern_name in METADATA_PATTERNS


def decide_action(
    severity: str,
    confidence: str,
    is_metadata: bool = False,
    has_value_finding: bool = False,
) -> str:
    """Map severity x confidence to an action: block, warn, or log.

    | Severity    | HIGH conf | MEDIUM conf | LOW conf |
    |-------------|-----------|-------------|----------|
    | critical    | block     | block       | warn     |
    | high        | block     | warn        | log      |
    | medium      | warn      | log         | log      |

    A metadata finding (topic word or field name) is downgraded from block to
    warn unless the file also contains a real VALUE match. Naming a column is
    not disclosing a record.
    """
    matrix = {
        ("critical", "high"): "block",
        ("critical", "medium"): "block",
        ("critical", "low"): "warn",
        ("high", "high"): "block",
        ("high", "medium"): "warn",
        ("high", "low"): "log",
        ("medium", "high"): "warn",
        ("medium", "medium"): "log",
        ("medium", "low"): "log",
    }
    action = matrix.get((severity, confidence), "block")

    if is_metadata and not has_value_finding and action == "block":
        action = "warn"

    return action


_ACTION_ORDER = {"block": 0, "warn": 1, "log": 2}


def worst_action(findings: list[dict]) -> str:
    """Return the most severe action across all findings in a file."""
    # Does this file contain any actual PII VALUE, as opposed to words about PII?
    has_value_finding = any(not _is_metadata(f["pattern_name"]) for f in findings)

    worst = "log"
    for f in findings:
        action = decide_action(
            f["severity"],
            f["confidence"],
            is_metadata=_is_metadata(f["pattern_name"]),
            has_value_finding=has_value_finding,
        )
        if _ACTION_ORDER[action] < _ACTION_ORDER[worst]:
            worst = action
    return worst


@dataclass(frozen=True)
class ReleasePolicy:
    floor: str = "high"
    strict: bool = False
    policy_error: str = ""

    def as_dict(self):
        result = {"floor": self.floor, "strict": self.strict}
        if self.policy_error:
            result["policy_error"] = self.policy_error
        return result


def resolve_policy(environ=None) -> ReleasePolicy:
    """Resolve once per invocation, never at import time or into cached findings."""
    environ = os.environ if environ is None else environ
    floor = environ.get("FERPA_GUARD_FLOOR", "").strip().lower() or "high"
    strict = bool(environ.get("FERPA_GUARD_STRICT"))
    if floor not in {"critical", "high", "medium"}:
        return ReleasePolicy("invalid", strict, "POLICY_CONFIG_INVALID")
    return ReleasePolicy(floor, strict)


def policy_error_finding() -> dict:
    return {"pattern_name": "POLICY_CONFIG_INVALID",
            "description": "POLICY_CONFIG_INVALID: configure FERPA_GUARD_FLOOR as critical, high or medium",
            "severity": "high", "confidence": "high", "count": 1,
            "header_only": False, "is_metadata": False}


def evaluate_policy(findings: list[dict], policy: ReleasePolicy) -> str:
    """Add release policy without changing the default matrix or findings."""
    if policy.policy_error:
        return "block"
    if not findings:
        return "allow"
    if policy.strict or any(f["pattern_name"] in {
            "SCAN_READER_UNAVAILABLE", "SCAN_INCOMPLETE", "POLICY_CONFIG_INVALID"} for f in findings):
        return "block"
    action = worst_action(findings)
    ranks = {"critical": 3, "high": 2, "medium": 1}
    for f in findings:
        if (not _is_metadata(f["pattern_name"]) and f["confidence"] == "high"
                and ranks.get(f["severity"], 0) >= ranks[policy.floor]):
            return "block"
    return action
